"""2026-10-26 원본(응답).xlsx 기준 jamovi 변환. 73건 전부 본 응답(시험응답 없음), 이전 분석 결과와 대조하지 않는다(사용자 지시)."""
import sys, json, hashlib
from pathlib import Path
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT=next(p for p in Path(__file__).resolve().parents if (p/'tools/survey_analysis/coding.py').exists())
sys.path.insert(0,str(ROOT))
from tools.survey_analysis import coding as c
OUT=Path(__file__).resolve().parent
PREFIX='261026_0631'
src=ROOT/c.SRC_REL
before=hashlib.sha256(src.read_bytes()).hexdigest()
c.EXCLUSIONS={}
r=c.build(src)
ts=r.raw.timestamp.map(lambda t:t.strftime('%Y-%m-%d %H:%M:%S') if hasattr(t,'strftime') else str(t))
r.analysis['include']=1
print('issues',len(r.issues));[print(' ',i) for i in r.issues]
assert not r.cross_fill_violations and not r.category_failures and not r.likert_failures
for item,col in {**c.ADOPT_COLS}.items():assert str(r.headers[col]).startswith(item[0]+'-'+item[1:]+')')
for item,col in c.NONADOPT_COLS.items():assert str(r.headers[col]).startswith(item[0]+'-'+item[1:]+')')
rename={'resp_id':'respondent_id','group':'D','A1_affil':'affiliation','A2_sex':'sex','A3_age':'age','A4_exp':'experience','A5_role':'role','A6_cost':'cost','A7_path':'A7','A8_when':'A8','A9_lag':'A9','W':'W_mean'}
cols=['resp_id','group','A1_affil','A2_sex','A3_age','A4_exp','A5_role','A6_cost','A7_path','A8_when','A9_lag']+c.LIKERT_ITEMS+list(c.DERIVED)+['sys_path']
all_data=r.analysis[cols].rename(columns=rename)
print('rows',len(all_data),'D counts',all_data.D.value_counts().to_dict())
same=None
meaning={'respondent_id':'비식별 응답 ID(원본 행 순서)','D':'시스템 도입 집단','affiliation':'소속 구분(A1)','sex':'성별(A2)','age':'만 나이(A3, 세)','experience':'건설현장 근무 경력(A4, 년)','role':'역할(A5)','cost':'총공사비 구간(A6)','Y_total':'안전관리 실효성 전체','Y1':'조기인지','Y2':'대응 신속성','Y3':'원격 파악','Y4':'순찰부담 경감','Y5':'위험 통제력','W_mean':'안전관리 예산 제약성','D_sub1':'인프라 범용성','D_sub2':'탐지·제어 안정성','D_sub3':'오탐 필터링','D_sub4':'알림 전파 즉각성','D_total':'시스템 특성 인식 전체','E_total':'경보 피로도 완화 인식','sys_path':'자기보고 시스템 인지경로 충족'}
for item in c.LIKERT_ITEMS+['A7','A8','A9']:
    meaning[item]=str(r.headers[c.ADOPT_COLS[item]])
cat_vars={'affiliation':'A1','sex':'A2','role':'A5','cost':'A6','A7':'A7','A8':'A8','A9':'A9'}
codebook=[]
for name in all_data:
    typ='연속 / Continuous';labels='';note=''
    if name=='respondent_id':typ='ID';note='문자열; 성명·현장명은 포함하지 않음'
    elif name=='D':typ='명목 / Nominal';labels='0=미도입; 1=도입';note='독립변수. D1~D8 문항과 구분'
    elif name in cat_vars:
        key=cat_vars[name];typ='순서 / Ordinal' if name in ('cost','A9') else '명목 / Nominal'
        labels='; '.join(f'{k}={v}' for k,v in c.CAT_LABELS[key].items())
        if name=='cost':note='ANCOVA에서는 순서형 숫자 공변량이 아닌 범주형 Fixed Factor로 투입'
    elif name in c.LIKERT_ITEMS:
        typ='순서 / Ordinal';labels='1=전혀 그렇지 않다; 2=그렇지 않다; 3=보통이다; 4=그렇다; 5=매우 그렇다'
        if name.startswith(('D','E')):note='도입 전용; 미도입은 설계상 비해당으로 빈칸'
    elif name=='sys_path':
        typ='명목 / Nominal';labels='0=아니오; 1=예'
        note='A7∈{3,4}이고 A8=3이면 1; 도입 집단에만 정의'
    else:
        oldname='W' if name=='W_mean' else name
        if oldname in c.DERIVED:
            items,adopt=c.DERIVED[oldname];note='평균('+','.join(items)+'); 구성 문항 전부 응답한 경우만 산출'
            if adopt:note+='; 미도입 빈칸'
    storage='Text' if name=='respondent_id' else ('Decimal' if name in ('age','experience') or name in [rename.get(k,k) for k in c.DERIVED] else 'Integer')
    codebook.append([name,meaning[name],storage,typ,labels,note])

def save(data,filename,kind):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='data'
    ws.append(list(data.columns))
    for row in data.itertuples(index=False,name=None):
        clean=[]
        for i,x in enumerate(row):
            if pd.isna(x):clean.append(None)
            elif i==0:clean.append(str(x))
            else:
                v=float(x);clean.append(int(v) if v.is_integer() else v)
        ws.append(clean)
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245678')
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=max(12,len(str(col[0].value))+2)
    cb=wb.create_sheet('codebook');cb.append(['variable','한국어 의미','데이터 형식','jamovi 측정수준','값 라벨','채점·사용 주의'])
    for row in codebook:cb.append(row)
    cb.freeze_panes='A2';cb.auto_filter.ref=cb.dimensions
    for k,width in {'A':20,'B':72,'C':18,'D':24,'E':90,'F':95}.items():cb.column_dimensions[k].width=width
    for row in cb:
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    info=wb.create_sheet('readme')
    notes=[['내용',kind],['가져올 시트','data (첫 번째 시트)'],['응답 수',len(data)],['변수 수',len(data.columns)],['원본 파일',str(c.SRC_REL)],['원본 SHA256',before],['생성 시각',datetime.now().isoformat()],['개인정보','성명·현장명·타임스탬프·자유기술 응답 제거'],['숫자 변환','응답은 숫자 셀; 비해당·결측은 빈칸, 0점 대체 없음'],['이전 자료 대조','수행하지 않음(2026-10-26 사용자 지시: 현재 원본을 기준으로 변환)'],['주의','XLSX에는 jamovi 값라벨·측정수준 메타데이터가 내장되지 않는다. codebook을 기준으로 가져온 뒤 변수 설정 확인'],['H1 설정','종속변수 Y_total; Fixed Factors D, affiliation, sex, role, cost; Covariates age, experience'],['H2 설정','W_mean을 분석 대상 전체 평균으로 중심화하여 W_c 생성 후 D×W_c 추가'],['표본','73건 전부 본 응답으로 분석(2026-10-26 사용자 판단: 시험응답 없음)'],['원고','본문·HWPX 수정 없음']]
    for row in notes:info.append(row)
    info.column_dimensions['A'].width=24;info.column_dimensions['B'].width=125
    for row in info:
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    path=OUT/filename;wb.save(path);return path

paths=[save(all_data,PREFIX+'_jamovi용_설문응답_73건.xlsx','jamovi 분석용: 73건 전부 포함')]
validation=[]
for path,expected in zip(paths,[all_data]):
    wb=openpyxl.load_workbook(path,data_only=True);ws=wb.worksheets[0]
    values=list(ws.values);assert values[0]==tuple(expected.columns)
    assert ws.max_row==len(expected)+1 and ws.max_column==len(expected.columns)
    for actual,exp in zip(values[1:],expected.itertuples(index=False,name=None)):
        for i,(x,y) in enumerate(zip(actual,exp)):
            if pd.isna(y):assert x is None
            elif i==0:assert x==str(y)
            else:assert isinstance(x,(int,float)) and abs(x-float(y))<1e-12
    non=[row for row in values[1:] if row[1]==0]
    assert len(values)-1==73
    for name in c.D_ITEMS+c.E_ITEMS+['D_total','E_total','D_sub1','D_sub2','D_sub3','D_sub4','sys_path']:
        j=list(expected.columns).index(name);assert all(row[j] is None for row in non)
    validation.append({'file':path.name,'rows':len(expected),'columns':len(expected.columns),'first_sheet':ws.title,'all_numeric_except_id':True,'missing_preserved':True,'roundtrip_exact':True})
    wb.close()
after=hashlib.sha256(src.read_bytes()).hexdigest();assert before==after
result={'source_sha_before':before,'source_sha_after':after,'validation':validation}
(OUT/(PREFIX+'_jamovi변환_검증.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False,indent=2))
