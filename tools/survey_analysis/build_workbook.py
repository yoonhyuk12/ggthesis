"""설문 응답 정리본 엑셀(01.docs/설문결과/261024_0850_설문응답_정리본.xlsx)과 coded_data.csv를 만든다.

저장소 루트에서 실행한다 (재실행 가능, 같은 입력이면 같은 결과).
    PYTHONIOENCODING=utf-8 python tools/survey_analysis/build_workbook.py [--no-recalc]

순서: 원본 SHA256 확인 → coding.build() → openpyxl로 7개 시트 작성 → CSV 저장
      → Excel COM 재계산(recalc_excel.ps1) → 재계산 값 검증(오류 셀·수식↔Python 대조) → 원본 SHA256 재확인.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from tools.survey_analysis import coding as cd  # noqa: E402

ROOT = cd.REPO_ROOT
OUT_XLSX_REL = Path("01.docs/설문결과/261024_0850_설문응답_정리본.xlsx")
OUT_CSV_REL = Path("01.docs/설문결과/analysis/coded_data.csv")
OUT_VERIFY_REL = Path("01.docs/설문결과/analysis/workbook_verification.json")
APPENDIX_REL = Path("01.docs/부록1_설문지_양식.md")
RECALC_PS1 = Path(__file__).resolve().parent / "recalc_excel.ps1"
KST = timezone(timedelta(hours=9))

FONT_NAME = "맑은 고딕"
F_INPUT = Font(name=FONT_NAME, size=10, color="0000FF")
F_FORMULA = Font(name=FONT_NAME, size=10, color="000000")
F_TEXT = Font(name=FONT_NAME, size=10, color="000000")
F_HEAD = Font(name=FONT_NAME, size=10, bold=True)
F_TITLE = Font(name=FONT_NAME, size=12, bold=True)
F_SECTION = Font(name=FONT_NAME, size=11, bold=True)
F_NOTE = Font(name=FONT_NAME, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", fgColor="D9E1F2")
FILL_SUB = PatternFill("solid", fgColor="F2F2F2")
FILL_CRIT = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
FORBIDDEN_FUNCS = ("XLOOKUP", "XMATCH", "FILTER(", "UNIQUE(", "SORT(", "SEQUENCE(", "_XLFN")
ERROR_STRINGS = ("#DIV/0!", "#VALUE!", "#NAME?", "#REF!", "#N/A", "#NUM!", "#NULL!")

# ---------------------------------------------------------------- 코딩자료 열 정의 (ASCII 키 → 시트 머리글)
CODED_HEADERS = {
    "resp_id": "resp_id", "src_row": "원본행", "group": "group",
    "A1_affil": "A1_소속", "A2_sex": "A2_성별", "A3_age": "A3_나이", "A4_exp": "A4_경력",
    "A5_role": "A5_역할", "A5_role_other": "A5_기타텍스트", "A6_cost": "A6_총공사비",
    "A7_path": "A7_인지경로", "A7_path_other": "A7_기타텍스트", "A8_when": "A8_인지시점", "A9_lag": "A9_전달소요",
    **{i: i for i in cd.LIKERT_ITEMS},
    "Y_total": "Y_total", "Y1": "Y1_조기인지", "Y2": "Y2_대응신속성", "Y3": "Y3_원격파악",
    "Y4": "Y4_순찰부담경감", "Y5": "Y5_위험통제력", "W": "W_예산제약",
    "D_sub1": "D_sub1_인프라범용성", "D_sub2": "D_sub2_탐지제어안정성", "D_sub3": "D_sub3_오탐필터링",
    "D_sub4": "D_sub4_알림전파즉각성", "D_total": "D_total", "E_total": "E_total",
    "sys_path": "sys_path", "include": "include", "exclude_reason": "exclude_reason",
}
CODED_KEYS = list(CODED_HEADERS)
CODED_COL = {k: i + 1 for i, k in enumerate(CODED_KEYS)}
CL = {k: get_column_letter(c) for k, c in CODED_COL.items()}   # ASCII 키 → 열 문자
N_ROWS = 73
LAST_ROW = N_ROWS + 1
CODED_SHEET = "코딩자료"
CS = f"'{CODED_SHEET}'!"
DB_RANGE = f"{CS}$A$1:${get_column_letter(len(CODED_KEYS))}${LAST_ROW}"

SUBSCALE_LABEL = {
    "Y_total": "안전관리 실효성 전체", "Y1": "조기인지", "Y2": "대응 신속성", "Y3": "원격 파악",
    "Y4": "순찰부담 경감", "Y5": "위험 통제력", "W": "안전관리 예산 제약성",
    "D_sub1": "인프라 범용성", "D_sub2": "탐지·제어 안정성", "D_sub3": "오탐 필터링",
    "D_sub4": "알림 전파 즉각성", "D_total": "시스템 특성 인식 전체", "E_total": "경보 피로도 완화 인식",
}
GROUPS = [("도입", 1), ("미도입", 0), ("전체", None)]


def colrange(key: str) -> str:
    return f"{CS}${CL[key]}$2:${CL[key]}${LAST_ROW}"


def text_width(s) -> float:
    s = "" if s is None else str(s)
    return sum(1.9 if ord(ch) > 0x1100 else 1.05 for ch in s)


def isblank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or v is pd.NA or v == ""


def clean(v):
    """pandas 값 → 셀 값 (NaN→None, 정수형 float→int)."""
    if isblank(v):
        return None
    if hasattr(v, "item"):
        v = v.item()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def style_header_row(ws, row: int, ncol: int, start: int = 1):
    for c in range(start, start + ncol):
        cell = ws.cell(row, c)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, CENTER, BORDER


def autofit(ws, min_w=6, max_w=45, skip_rows=()):
    widths: dict[int, float] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.row in skip_rows or cell.value is None:
                continue
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                w = 9
            elif isinstance(v, datetime):
                w = 19
            elif isinstance(v, float):
                w = 8
            else:
                w = max(text_width(x) for x in str(v).split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 0), w)
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = max(min_w, min(max_w, w + 2))


# ---------------------------------------------------------------- 부록1 문항 전문
def appendix_items() -> dict[str, str]:
    text = (ROOT / APPENDIX_REL).read_text(encoding="utf-8")
    items: dict[str, str] = {}
    for m in re.finditer(r"^\*\*A-(\d)\)\*\*\s*(.+)$", text, flags=re.M):
        items[f"A{m.group(1)}"] = m.group(2).strip()
    for m in re.finditer(r"^\|\s*([B-E])-(\d)\s*\|\s*([^|]+?)\s*\|", text, flags=re.M):
        items.setdefault(f"{m.group(1)}{m.group(2)}", m.group(3).strip())
    return items


def xlsx_item_text(headers: dict[int, str], item: str) -> str:
    h = str(headers[cd.ADOPT_COLS[item]])
    return re.sub(r"^[A-E]-\d\)\s*", "", h).strip()


# ---------------------------------------------------------------- 시트 작성
def write_table(ws, df_rows: list[list], header: list[str], formula_cols: set[int] = frozenset(),
                text_cols: set[int] = frozenset(), date_cols: set[int] = frozenset(), freeze: str = "B2"):
    # freeze_panes는 시트당 한 번만 지정한다(두 번 지정하면 중복 selection 요소로 Excel이 파일을 열지 못한다)
    for c, h in enumerate(header, 1):
        ws.cell(1, c, h)
    style_header_row(ws, 1, len(header))
    for r, rec in enumerate(df_rows, 2):
        for c, v in enumerate(rec, 1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
            if c in formula_cols:
                cell.font = F_FORMULA
                cell.number_format = "0.000"
            else:
                cell.font = F_INPUT
            if c in date_cols:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if c in text_cols:
                cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(df_rows) + 1}"


def sheet_raw(wb, res: cd.BuildResult, item_text: dict[str, str]):
    ws = wb.create_sheet("원자료_통합")
    header = ["resp_id", "원본행", "타임스탬프", "집단", "group", "현장명", "성명"]
    header += [f"{i[0]}-{i[1:]}" for i in cd.ITEMS] + ["include", "exclude_reason"]
    rows = []
    for _, r in res.raw.iterrows():
        ts = r["timestamp"]
        ts = ts.to_pydatetime() if hasattr(ts, "to_pydatetime") else ts
        rec = [r.resp_id, int(r.src_row), ts, r.group_text, clean(r.group), clean(r.site), clean(r["name"])]
        rec += [clean(r[i]) for i in cd.ITEMS] + [int(r.include), clean(r.exclude_reason)]
        rows.append(rec)
    write_table(ws, rows, header, date_cols={3})
    for c, item in enumerate(cd.ITEMS, 8):
        ws.cell(1, c).comment = Comment(f"{item[0]}-{item[1:]}) {item_text[item]}", "build_workbook")
    autofit(ws, max_w=30)
    return ws


def derived_formula(key: str, r: int) -> str:
    items, adopt_only = cd.DERIVED[key]
    rng = f"{CL[items[0]]}{r}:{CL[items[-1]]}{r}"
    core = f'IF(COUNT({rng})<{len(items)},"",AVERAGE({rng}))'
    if adopt_only:
        return f'=IF(${CL["group"]}{r}<>1,"",{core})'
    return f"={core}"


def sys_path_formula(r: int) -> str:
    g, a7, a8 = f'${CL["group"]}{r}', f'{CL["A7_path"]}{r}', f'{CL["A8_when"]}{r}'
    return (f'=IF({g}<>1,"",IF(OR({a7}="",{a8}=""),"",'
            f'IF(AND(OR({a7}=3,{a7}=4),{a8}=3),1,0)))')


def sheet_coded(wb, res: cd.BuildResult):
    ws = wb.create_sheet(CODED_SHEET)
    header = [CODED_HEADERS[k] for k in CODED_KEYS]
    formula_keys = set(cd.DERIVED) | {"sys_path"}
    rows = []
    for idx, (_, r) in enumerate(res.coded.iterrows()):
        excel_row = idx + 2
        rec = []
        for k in CODED_KEYS:
            if k in cd.DERIVED:
                rec.append(derived_formula(k, excel_row))
            elif k == "sys_path":
                rec.append(sys_path_formula(excel_row))
            else:
                rec.append(clean(r[k]))
        rows.append(rec)
    write_table(ws, rows, header, formula_cols={CODED_COL[k] for k in formula_keys})
    for k in ("sys_path",):
        for r in range(2, LAST_ROW + 1):
            ws.cell(r, CODED_COL[k]).number_format = "0"
    autofit(ws, max_w=28)
    for k in formula_keys:
        ws.column_dimensions[CL[k]].width = max(9, text_width(CODED_HEADERS[k]) + 2)
    return ws


def analysis_frame(res: cd.BuildResult) -> pd.DataFrame:
    a = res.analysis.copy()
    float_cols = [c for c in cd.DERIVED]
    for c in a.columns:
        if c in float_cols or c in ("resp_id", "A5_role_other", "A7_path_other", "exclude_reason"):
            continue
        a[c] = pd.array([None if isblank(v) else int(v) for v in a[c]], dtype="Int64")
    for c in float_cols:
        a[c] = pd.to_numeric(a[c], errors="coerce")
    return a


def sheet_analysis(wb, a: pd.DataFrame):
    ws = wb.create_sheet("분석자료")
    rows = [[clean(v) for v in rec] for rec in a.itertuples(index=False, name=None)]
    write_table(ws, rows, list(a.columns))
    for c, name in enumerate(a.columns, 1):
        if name in cd.DERIVED:
            for r in range(2, len(rows) + 2):
                ws.cell(r, c).number_format = "0.000"
    autofit(ws, max_w=24)
    return ws


ROLE_TEXT = {
    "resp_id": ("식별자", "응답 1건 식별자(원본 행 순서 R001∼R073)"),
    "group": ("독립변수(D)", "<표 4-3> 독립변수. 자기보고 분기값(명부 대조 전)"),
    "A1_affil": ("조정변수: 범주형 통제요인", "기준범주=시공사(04장 제5절 제2항)"),
    "A2_sex": ("조정변수: 범주형 통제요인", "기준범주=남성"),
    "A3_age": ("조정변수: 연속 공변량", "연속값 그대로 투입"),
    "A4_exp": ("조정변수: 연속 공변량", "연속값 그대로 투입"),
    "A5_role": ("조정변수: 범주형 통제요인", "기준범주=공사감독, 4수준"),
    "A5_role_other": ("참고(기타 원문)", "A5=4일 때 원문"),
    "A6_cost": ("조정변수: 범주형 통제요인", "기준범주=20억 원 미만. 조사용 4구간(등간 아님)"),
    "A7_path": ("참고 변수: 현장 운영 특성", "운용 실태 확인 지표(가설 검정 미투입)"),
    "A7_path_other": ("참고(기타 원문)", "A7=5일 때 원문"),
    "A8_when": ("참고 변수: 현장 운영 특성", "운용 실태 확인 지표(가설 검정 미투입)"),
    "A9_lag": ("참고 변수: 현장 운영 특성", "과정 지표 — <표 4-6> 조합 판정에 사용"),
}


def coding_rule(key: str) -> str:
    base = key.split("_")[0] if key[:2] in {f"A{i}" for i in range(1, 10)} else key
    if key in ("A5_role_other", "A7_path_other"):
        return "원문 텍스트(strip)"
    if base in cd.CAT_MAPS:
        rule = ", ".join(f"{v}={t}" for v, t in sorted(cd.CAT_LABELS[base].items()))
        if base == "A5":
            rule += " (그 외 텍스트=4, '.'은 무응답→결측)"
        if base == "A7":
            rule += " (그 외 텍스트=5)"
        return rule
    if base in ("A3", "A4"):
        return "숫자 그대로(연속)"
    if key in cd.LIKERT_ITEMS:
        return "① 전혀 그렇지 않다=1 … ⑤ 매우 그렇다=5 (동그라미 숫자로 판정, 그 외=결측)"
    if key == "group":
        return f"1={cd.BRANCH_ADOPT}, 0={cd.BRANCH_NONADOPT}"
    if key in cd.DERIVED:
        items, adopt_only = cd.DERIVED[key]
        s = f"mean({', '.join(items)}) — {len(items)}문항 전부 응답 시에만 산출, 아니면 빈칸"
        return s + ("; 도입 집단만(미도입 빈칸)" if adopt_only else "")
    if key == "sys_path":
        return "도입 집단: A7∈{3,4} 그리고 A8=3 → 1, 그 외 0; 미도입 빈칸. A7=5(기타)는 0 [확정 필요]"
    if key == "include":
        return "1=분석 포함, 0=제외(시험 응답 2건)"
    if key == "exclude_reason":
        return "include=0의 사유"
    return ""


def sheet_codebook(wb, item_text: dict[str, str]):
    ws = wb.create_sheet("코드북")
    header = ["변수명(분석자료·CSV)", "코딩자료 열 이름", "블록", "문항 전문(부록1)", "코딩 규칙", "하위요인", "분석 역할(<표 4-3>)", "비고"]
    sub_of = {}
    for k, (items, _) in cd.DERIVED.items():
        if k in ("Y_total", "D_total", "E_total", "W"):
            continue
        for i in items:
            sub_of[i] = f"{k} {SUBSCALE_LABEL[k]}"
    rows = []
    for key in cd.ANALYSIS_COLUMNS:
        coded_name = CODED_HEADERS.get(key, key)
        base = key.split("_")[0]
        if key in cd.LIKERT_ITEMS:
            block = {"B": "Part B 안전관리 실효성", "C": "Part C 예산 제약성", "D": "Part D 시스템 특성(도입 전용)",
                     "E": "Part E 경보 피로도(도입 전용)"}[key[0]]
            text = item_text.get(key, "")
            if key[0] == "B":
                role, note = "종속변수(Y) 구성 문항", ""
            elif key[0] == "C":
                role, note = "상호작용 변수(W) 구성 문항", ""
            elif key[0] == "D":
                role, note = "참고 변수: 시스템 특성 인식(도입 집단)", "가설 검정 미투입, 미도입 빈칸"
            else:
                role, note = "참고 변수: 경보 피로도 완화 인식(도입 집단)", "가설 검정 미투입, 미도입 빈칸"
            sub = sub_of.get(key, "E_total" if key[0] == "E" else "")
        elif base in {f"A{i}" for i in range(1, 10)}:
            block = "Part A 일반 사항·현장 운영 특성"
            text = item_text.get(base, "") if not key.endswith("_other") else f"{base} 기타(  ) 원문"
            role, note = ROLE_TEXT[key]
            sub = ""
        elif key in cd.DERIVED:
            block = "파생 변수"
            items, _ = cd.DERIVED[key]
            text = f"{SUBSCALE_LABEL[key]} ({'·'.join(i[0] + '-' + i[1:] for i in items)})"
            sub = SUBSCALE_LABEL[key]
            if key == "Y_total":
                role, note = "종속변수(Y) — H1·H2", "영역 가중 12.5·12.5·25·25·25%"
            elif key.startswith("Y"):
                role, note = "종속변수 하위요인 — 영역별 분석(<표 5-7>, Holm)", ""
            elif key == "W":
                role, note = "상호작용 변수(W) — H2", "분석 시 완전사례 평균으로 중심화(Task B)"
            else:
                role, note = "참고 변수(도입 집단 기술통계)", "가설 검정 미투입"
        elif key == "sys_path":
            block, text, sub = "파생 변수", "시스템 경로 작동(부록1·05장 <표 5-1> 주 규칙)", ""
            role, note = "민감도 분석용(<표 5-6> sys_path=0 제외 재추정)", "[확정 필요: A-7 ⑤ 판정 기준]"
        elif key in ("include", "exclude_reason"):
            block, text, sub = "관리 변수", "분석 포함 여부" if key == "include" else "제외 사유", ""
            role, note = "표본 선택(jamovi에서 include=1 필터)", ""
        else:
            block, text, sub = "관리 변수", "", ""
            role, note = ROLE_TEXT.get(key, ("", ""))
        rows.append([key, coded_name, block, text, coding_rule(key), sub, role, note])
    write_table(ws, rows, header)
    for r in range(2, len(rows) + 2):
        for c in range(1, len(header) + 1):
            ws.cell(r, c).font = F_TEXT
            ws.cell(r, c).alignment = WRAP
    for c, w in zip("ABCDEFGH", (16, 22, 24, 60, 50, 22, 36, 34)):
        ws.column_dimensions[c].width = w
    return ws


# ---------------------------------------------------------------- 빈도_기술통계 (수식) + 기대값
class StatsSheet:
    def __init__(self, ws, coded: pd.DataFrame):
        self.ws = ws
        self.r = 1
        self.expected: dict[str, object] = {}
        self.inc = coded[coded.include == 1]
        # D함수 조건 범위 (K:L)
        ws["K1"] = "D함수 조건 범위(수정 금지)"
        ws["K1"].font = F_HEAD
        self.crit = {}
        spots = {1: (2, True), 0: (5, True), None: (8, False)}
        for g, (row, two) in spots.items():
            label = {1: "도입", 0: "미도입", None: "전체"}[g]
            ws.cell(row, 10, label).font = F_HEAD
            if two:
                ws.cell(row, 11, "group"); ws.cell(row, 12, "include")
                ws.cell(row + 1, 11, g); ws.cell(row + 1, 12, 1)
                self.crit[g] = f"$K${row}:$L${row + 1}"
                cells = [(row, 11), (row, 12), (row + 1, 11), (row + 1, 12)]
            else:
                ws.cell(row, 11, "include"); ws.cell(row + 1, 11, 1)
                self.crit[g] = f"$K${row}:$K${row + 1}"
                cells = [(row, 11), (row + 1, 11)]
            for rr, cc in cells:
                cell = ws.cell(rr, cc)
                cell.fill, cell.border = FILL_CRIT, BORDER
                cell.font = F_INPUT if rr == row + 1 else F_HEAD
        ws.column_dimensions["J"].width = 8
        ws.column_dimensions["K"].width = 10
        ws.column_dimensions["L"].width = 10

    def sub(self, g):
        return self.inc if g is None else self.inc[self.inc.group == g]

    def put(self, row, col, formula, expected, fmt=None):
        cell = self.ws.cell(row, col, formula)
        cell.font, cell.border = F_FORMULA, BORDER
        if fmt:
            cell.number_format = fmt
        self.expected[cell.coordinate] = expected

    def label(self, row, col, v, bold=False, fill=None):
        cell = self.ws.cell(row, col, v)
        cell.font = F_HEAD if bold else F_TEXT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            cell.fill = fill

    def title(self, text, note=None):
        self.r += 1
        self.ws.cell(self.r, 1, text).font = F_SECTION
        self.r += 1
        if note:
            self.ws.cell(self.r, 1, note).font = F_NOTE
            self.r += 1

    def header(self, labels):
        for c, h in enumerate(labels, 1):
            self.ws.cell(self.r, c, h)
        style_header_row(self.ws, self.r, len(labels))
        self.r += 1

    # COUNTIFS 조건 문자열
    def cond(self, g, extra: list[str]):
        parts = []
        if g is not None:
            parts += [colrange("group"), str(g)]
        parts += [colrange("include"), "1"] + extra
        return ",".join(parts)

    def categorical(self, title_text, items: list[tuple[str, str, dict[int, str]]], note=None):
        self.title(title_text, note)
        self.header(["구분", "항목", "도입 n", "도입 %", "미도입 n", "미도입 %", "전체 n", "전체 %"])
        for var_label, key, labels in items:
            valid_row = self.r + len(labels)
            first = True
            for code, lab in labels.items():
                self.label(self.r, 1, var_label if first else None)
                self.label(self.r, 2, lab)
                first = False
                for j, (_, g) in enumerate(GROUPS):
                    s = self.sub(g)
                    n = int((s[key] == code).sum())
                    valid = int(s[key].notna().sum())
                    ncell = get_column_letter(3 + 2 * j)
                    self.put(self.r, 3 + 2 * j, f"=COUNTIFS({self.cond(g, [colrange(key), str(code)])})", n, "0")
                    self.put(self.r, 4 + 2 * j, f'=IFERROR({ncell}{self.r}/{ncell}{valid_row},"")',
                             (n / valid) if valid else None, "0.0%")
                self.r += 1
            self.label(self.r, 1, None, fill=FILL_SUB)
            self.label(self.r, 2, "유효 응답(분모)", fill=FILL_SUB)
            for j, (_, g) in enumerate(GROUPS):
                s = self.sub(g)
                self.put(self.r, 3 + 2 * j, f'=COUNTIFS({self.cond(g, [colrange(key), chr(34) + "<>" + chr(34)])})',
                         int(s[key].notna().sum()), "0")
                total = len(s)
                self.put(self.r, 4 + 2 * j, f"=COUNTIFS({self.cond(g, [])})-{get_column_letter(3 + 2 * j)}{self.r}",
                         total - int(s[key].notna().sum()), '"결측 "0')
                self.ws.cell(self.r, 3 + 2 * j).fill = FILL_SUB
                self.ws.cell(self.r, 4 + 2 * j).fill = FILL_SUB
            self.r += 1

    def total_row(self):
        self.label(self.r, 1, "합계(include=1)", bold=True)
        self.label(self.r, 2, None)
        for j, (_, g) in enumerate(GROUPS):
            self.put(self.r, 3 + 2 * j, f"=COUNTIFS({self.cond(g, [])})", len(self.sub(g)), "0")
            self.label(self.r, 4 + 2 * j, None)
        self.r += 1

    def descriptive(self, title_text, keys: list[tuple[str, str]], groups=GROUPS, note=None, minmax=True):
        self.title(title_text, note)
        head = ["변수", "집단", "n", "평균(M)", "표준편차(SD)"] + (["최솟값", "최댓값"] if minmax else [])
        self.header(head)
        for var_label, key in keys:
            field = CODED_HEADERS[key]
            first = True
            for glabel, g in groups:
                s = self.sub(g)[key].dropna().astype(float)
                self.label(self.r, 1, var_label if first else None)
                self.label(self.r, 2, glabel)
                first = False
                crit = self.crit[g]
                fn = lambda f: f'=IFERROR({f}({DB_RANGE},"{field}",{crit}),"")'  # noqa: E731
                self.put(self.r, 3, fn("DCOUNT"), len(s), "0")
                self.put(self.r, 4, fn("DAVERAGE"), s.mean() if len(s) else None, "0.00")
                self.put(self.r, 5, fn("DSTDEV"), s.std(ddof=1) if len(s) > 1 else None, "0.00")
                if minmax:
                    self.put(self.r, 6, fn("DMIN"), s.min() if len(s) else None, "General")
                    self.put(self.r, 7, fn("DMAX"), s.max() if len(s) else None, "General")
                self.r += 1

    def de_items(self):
        self.title("표 5-3 패널 B — D·E 문항별 응답 상태와 기술통계(도입 집단)",
                   "일반 무응답 = 도입 집단 include=1 인원 − 유효 응답 수. 경험 부재 비해당 선택지는 설문에 없음(—).")
        self.header(["문항", "집단", "유효 응답 수", "일반 무응답 수", "평균(M)", "표준편차(SD)"])
        crit = self.crit[1]
        s_all = self.sub(1)
        for key in cd.D_ITEMS + cd.E_ITEMS:
            s = s_all[key].dropna().astype(float)
            self.label(self.r, 1, f"{key[0]}-{key[1:]}")
            self.label(self.r, 2, "도입")
            self.put(self.r, 3, f'=DCOUNT({DB_RANGE},"{key}",{crit})', len(s), "0")
            self.put(self.r, 4, f"=COUNTIFS({self.cond(1, [])})-C{self.r}", len(s_all) - len(s), "0")
            self.put(self.r, 5, f'=IFERROR(DAVERAGE({DB_RANGE},"{key}",{crit}),"")', s.mean(), "0.00")
            self.put(self.r, 6, f'=IFERROR(DSTDEV({DB_RANGE},"{key}",{crit}),"")', s.std(ddof=1), "0.00")
            self.r += 1

    def sys_path_block(self):
        self.title("도입 집단 시스템 경로 작동(sys_path) 집계", "sys_path=1: A-7이 ③·④이고 A-8이 ③. A-7 ⑤(기타)는 0으로 집계 [확정 필요].")
        self.header(["구분", "항목", "도입 n"])
        s = self.sub(1)
        rows = [("sys_path=1 (경로 작동)", [colrange("sys_path"), "1"], int((s.sys_path == 1).sum())),
                ("sys_path=0 (경로 미작동)", [colrange("sys_path"), "0"], int((s.sys_path == 0).sum())),
                ("그중 A-7 ⑤ 기타", [colrange("A7_path"), "5"], int((s.A7_path == 5).sum()))]
        for lab, extra, exp in rows:
            self.label(self.r, 1, "시스템 경로")
            self.label(self.r, 2, lab)
            self.put(self.r, 3, f"=COUNTIFS({self.cond(1, extra)})", exp, "0")
            self.r += 1


def sheet_stats(wb, coded: pd.DataFrame):
    ws = wb.create_sheet("빈도_기술통계")
    st = StatsSheet(ws, coded)
    ws["A1"] = "빈도·기술통계 — 코딩자료 참조 수식, include=1만 집계 (집단: 명부 대조 전 자기보고 분기값)"
    ws["A1"].font = F_TITLE
    L = cd.CAT_LABELS
    st.categorical("표 5-1 패널 A — 범주형 특성의 빈도와 비율",
                   [("소속 구분", "A1_affil", L["A1"]), ("성별", "A2_sex", L["A2"]),
                    ("역할", "A5_role", L["A5"]), ("총공사비 구간", "A6_cost", L["A6"])],
                   "비율 분모 = 항목별 유효 응답. 각 변수 아래 회색 행: 유효 응답 n과 결측 수.")
    st.total_row()
    st.descriptive("표 5-1 패널 B — 만 나이·건설현장 경력",
                   [("만 나이(세)", "A3_age"), ("건설현장 경력(년)", "A4_exp")],
                   note="중위값은 수식 제약(MEDIAN 조건부 배열식 미사용)으로 이 시트에서 산출하지 않는다 — Task B 분석 결과 참조.")
    labels7 = {k: f"{'①②③④⑤'[k - 1]} {v}" for k, v in L["A7"].items()}
    labels8 = {k: f"{'①②③④'[k - 1]} {v}" for k, v in L["A8"].items()}
    labels9 = {k: f"{'①②③④'[k - 1]} {v}" for k, v in L["A9"].items()}
    st.categorical("표 5-1 패널 C — 현장 운영 특성(A-7∼A-9)",
                   [("A-7 담당자 부재 시 위험 인지 경로", "A7_path", labels7),
                    ("A-8 순찰 간 위험 인지 시점", "A8_when", labels8),
                    ("A-9 인지 후 지시 전달 소요시간", "A9_lag", labels9)],
                   "Fisher 정확검정 p는 Task B 분석에서 산출.")
    st.sys_path_block()
    st.descriptive("표 5-3 패널 A — 주요 변수의 기술통계(집단별)",
                   [(SUBSCALE_LABEL[k], k) for k in ("Y_total", "Y1", "Y2", "Y3", "Y4", "Y5", "W")],
                   note="왜도·첨도는 Task B 분석 결과 참조. n은 해당 척도 전 문항 응답자 수.")
    st.descriptive("표 5-3 패널 A(계속) — 도입 집단 참고 척도",
                   [(SUBSCALE_LABEL[k], k) for k in ("E_total", "D_total", "D_sub1", "D_sub2", "D_sub3", "D_sub4")],
                   groups=[("도입", 1)])
    st.de_items()
    st.descriptive("부가 — B·C 문항별 기술통계(집단별)",
                   [(f"{k[0]}-{k[1:]}", k) for k in cd.B_ITEMS + cd.C_ITEMS], minmax=True)
    ws.freeze_panes = "A2"
    for c, w in zip("ABCDEFGH", (30, 34, 10, 10, 12, 10, 10, 10)):
        ws.column_dimensions[c].width = w
    return st


# ---------------------------------------------------------------- 검토메모
def review_items(res: cd.BuildResult, item_text: dict[str, str]) -> list[list]:
    a, raw = res.coded, res.raw
    inc = a[a.include == 1]
    ad, na = inc[inc.group == 1], inc[inc.group == 0]
    memos: list[list] = []

    def add(cat, rid, row, detail, action, status="기록"):
        memos.append([len(memos) + 1, cat, rid, row, detail, action, status])

    # 0.5 사전 발견 특성
    sex = pd.crosstab(inc.group, inc.A2_sex)
    add("자료 특성(0.5)", "—", "—",
        f"성별: 도입 집단 {len(ad)}명 전원 남성(여성 0), 여성은 미도입에만 {int((na.A2_sex == 2).sum())}명.",
        "성별 더미는 미도입 내부에서만 식별됨 → 집단과 성별이 부분 교락. Task B에서 계수 해석 시 유의.")
    xt = pd.crosstab([inc.group, inc.A1_affil], inc.A5_role)
    cells = []
    for (g, af), row in xt.iterrows():
        cells.append(f"{'도입' if g == 1 else '미도입'}·{cd.CAT_LABELS['A1'][af][:3]}: " +
                     ", ".join(f"{cd.CAT_LABELS['A5'][int(k)]} {int(v)}" for k, v in row.items() if v))
    add("자료 특성(0.5)", "—", "—",
        "소속×역할 교차(include=1): " + " / ".join(cells) + ". 발주청 응답자는 대부분 공사감독.",
        "소속–역할 중복으로 설계행렬 순위 결함 가능 → Task B에서 rank(X) 확인 필수.")
    contractor_sup = inc[(inc.A1_affil == 1) & (inc.A5_role == 1)]
    if len(contractor_sup):
        add("자료 특성(추가)", ", ".join(contractor_sup.resp_id), ", ".join(map(str, contractor_sup.src_row)),
            f"소속=시공사인데 역할=공사감독 {len(contractor_sup)}건(보기상 공사감독은 발주청 역할에 가까움).",
            "응답 그대로 유지. 소속 또는 역할 오기 가능성 — 명부 대조 시 확인.", "확인 필요")
    owner_safety = inc[(inc.A1_affil == 2) & (inc.A5_role == 2)]
    if len(owner_safety):
        add("자료 특성(추가)", ", ".join(owner_safety.resp_id), ", ".join(map(str, owner_safety.src_row)),
            f"소속=발주청인데 역할=안전관리자 {len(owner_safety)}건.", "응답 그대로 유지(참고).")
    named = raw[raw.site.notna() | raw["name"].notna()]
    add("자료 특성(0.5)", ", ".join(named.resp_id), ", ".join(map(str, named.src_row)),
        f"현장명·성명은 {len(named)}건에만 입력(시험 응답 2건 포함). 나머지 {len(raw) - len(named)}건은 둘 다 빈칸.",
        "중복 제출은 타임스탬프·응답 벡터 동일성으로만 확인 가능(아래 항목). 분석자료·CSV에서는 성명·현장명 제거.")
    add("자료 특성(0.5)", "—", "—",
        "집단 구분은 첫 질문의 자기보고 분기값이며 조사 명부와의 대조는 아직 하지 않았다.",
        "group은 '명부 대조 전 자기보고 기준'. 명부 확보 후 불일치 건을 재검토(자기보고만으로 재분류하지 않음).", "확인 필요")

    # 중복·근접 제출
    keys = ["group", "A1_affil", "A2_sex", "A3_age", "A4_exp", "A5_role", "A6_cost", "A7_path", "A8_when", "A9_lag"] + cd.LIKERT_ITEMS
    dup = a[a.duplicated(subset=keys, keep=False)]
    add("중복 의심", ", ".join(dup.resp_id) if len(dup) else "—", "—",
        f"전 문항(A·B·C·D·E 코딩값) 응답 벡터가 완전히 같은 쌍: {len(dup)}건.", "중복 제출 근거 없음." if not len(dup) else "검토 필요")
    ts = raw.timestamp
    near = []
    for i in range(1, len(raw)):
        sec = (ts.iloc[i] - ts.iloc[i - 1]).total_seconds()
        if sec < 120:
            x, y = a.iloc[i - 1], a.iloc[i]
            ndiff = sum(1 for k in keys if not (isblank(x[k]) and isblank(y[k])) and x[k] != y[k])
            near.append((raw.resp_id.iloc[i - 1], raw.resp_id.iloc[i], int(sec), ndiff))
    if near:
        tight = [n for n in near if n[2] <= 20]
        add("중복 의심(근접 제출)", "; ".join(f"{p}–{q}" for p, q, _, _ in tight),
            "—",
            f"직전 응답과 2분 이내 제출 {len(near)}쌍. 20초 이내: " +
            "; ".join(f"{p}–{q} {s}초 차, 코딩값 {d}개 다름" for p, q, s, d in tight) +
            f". 2분 이내 쌍의 서로 다른 변수 수 최소 {min(n[3] for n in near)}개.",
            "근접 제출이지만 응답 벡터가 다르므로 중복으로 보지 않음(현장 단체 안내 시 동시 응답으로 해석). 계획서의 행 14·15 = R013–R014.")
    # 제외
    for _, r in a[a.include == 0].iterrows():
        rr = raw[raw.resp_id == r.resp_id].iloc[0]
        vals = sorted({int(r[k]) for k in cd.LIKERT_ITEMS if not isblank(r[k])})
        add("제외(include=0)", r.resp_id, int(r.src_row),
            f"{rr.timestamp:%Y-%m-%d %H:%M:%S}, 현장명 '{rr.site}', {'도입' if r.group == 1 else '미도입'} 분기, 리커트 전 문항 값 {vals}.",
            f"include=0, exclude_reason='{r.exclude_reason}'. → 유효 표본에서 제외.", "사용자 확인 필요(계획 결정 1)")
    # 무응답 표기
    for i in res.issues:
        if i.category == "기타 자리 무응답 표기":
            add("결측(무응답 표기)", i.resp_id, i.src_row, f"{i.detail} — 역할 보기 외 '.' 한 글자.",
                "A5_role 결측. H1 완전사례에서 탈락. (계획 결정 2: 결측 vs 기타)", "사용자 확인 필요(계획 결정 2)")
    # 기타 텍스트
    o5 = a[a.A5_role == 4]
    add("기타 응답(A-5)", ", ".join(o5.resp_id), ", ".join(map(str, o5.src_row)),
        "A-5 기타 원문: " + "; ".join(f"{r.resp_id}={r.A5_role_other}" for _, r in o5.iterrows()),
        "A5_role=4, 원문은 A5_기타텍스트 보존.")
    o7 = a[a.A7_path == 5]
    add("기타 응답(A-7)", ", ".join(o7.resp_id), ", ".join(map(str, o7.src_row)),
        "A-7 기타 원문: " + "; ".join(f"{r.resp_id}({'도입' if r.group == 1 else '미도입'})={r.A7_path_other}" for _, r in o7.iterrows()),
        "A7_path=5. 도입 집단 해당 건은 sys_path=0으로 두었으며 판정 기준 [확정 필요].", "확인 필요")
    # sys_path
    sp1 = ad[ad.sys_path == 1]
    xt78 = pd.crosstab(ad.A7_path, ad.A8_when)
    add("운용 실태(sys_path)", ", ".join(sp1.resp_id), ", ".join(map(str, sp1.src_row)),
        f"도입 {len(ad)}명 중 sys_path=1은 {len(sp1)}명, 0은 {int((ad.sys_path == 0).sum())}명. "
        f"A-7 분포(도입): " + ", ".join(f"{cd.CAT_LABELS['A7'][int(k)]} {int(v)}" for k, v in ad.A7_path.value_counts().sort_index().items()) +
        f". sys_path=1의 구성: A-7=③·A-8=③ {int(((sp1.A7_path == 3) & (sp1.A8_when == 3)).sum())}건, A-7=④·A-8=③ {int(((sp1.A7_path == 4) & (sp1.A8_when == 3)).sum())}건.",
        "sys_path=0 제외 민감도 분석의 부분집단이 매우 작다(도입 n=%d) → Task B 민감도 분석의 추정 안정성 주의." % len(sp1),
        "확인 필요")
    na_alert = na[(na.A8_when == 3) | (na.A7_path == 4)]
    if len(na_alert):
        add("운용 실태(미도입)", ", ".join(na_alert.resp_id), ", ".join(map(str, na_alert.src_row)),
            "미도입 응답인데 A-8=③(발생 직후 알림으로) 또는 A-7=④(자동 알림): " +
            "; ".join(f"{r.resp_id} A7={int(r.A7_path)}, A8={int(r.A8_when)}" for _, r in na_alert.iterrows()),
            "다른 알림 수단 사용 가능성 또는 분기 오선택 가능성 — 명부 대조 대상. 응답 그대로 유지.", "확인 필요")
    # 일렬응답
    for _, r in inc.iterrows():
        items = cd.B_ITEMS + cd.C_ITEMS + (cd.D_ITEMS + cd.E_ITEMS if r.group == 1 else [])
        vals = {r[k] for k in items}
        if len(vals) == 1:
            add("일렬응답", r.resp_id, int(r.src_row),
                f"{'도입' if r.group == 1 else '미도입'} 응답자의 리커트 전 문항({len(items)}개)이 모두 {int(list(vals)[0])}.",
                "원고 규칙(05장 제1절: 일렬응답만으로 자동 제외하지 않음)에 따라 포함 유지. 민감도 점검 후보.")
    # 희소 범주
    sparse = []
    for key, lab in (("A6_cost", "A6"), ("A5_role", "A5"), ("A2_sex", "A2"), ("A1_affil", "A1")):
        for g, gl in ((1, "도입"), (0, "미도입")):
            vc = inc[inc.group == g][key].value_counts()
            for code, label in cd.CAT_LABELS[lab].items():
                n = int(vc.get(code, 0))
                if n < 5:
                    sparse.append(f"{gl} {label} {n}")
    add("희소 범주(n<5)", "—", "—", "집단별 n<5 범주: " + "; ".join(sparse) + ".",
        "<표 5-4> 균형 진단에서 희소범주 표시. 더미 계수 불안정 가능.")
    # 검증 결과
    add("검증", "—", "—", f"교차 채움 위반 {len(res.cross_fill_violations)}건(도입 행 38∼58열, 미도입 행 5∼37열 모두 빈칸).", "이상 없음")
    add("검증", "—", "—",
        f"리커트 매핑 실패 {len(res.likert_failures)}건, 범주 매핑 실패(기타 제외) {len(res.category_failures)}건, 숫자 변환 실패 0건.",
        "이상 없음" if not (res.likert_failures or res.category_failures) else "결측 처리")
    miss = {k: int(inc[k].isna().sum()) for k in ["A1_affil", "A2_sex", "A3_age", "A4_exp", "A5_role", "A6_cost",
                                                   "A7_path", "A8_when", "A9_lag"] + cd.B_ITEMS + cd.C_ITEMS}
    miss_d = {k: int(ad[k].isna().sum()) for k in cd.D_ITEMS + cd.E_ITEMS}
    nz = {k: v for k, v in {**miss, **miss_d}.items() if v}
    add("결측 현황", "—", "—",
        "include=1 기준 문항 결측: " + (", ".join(f"{k} {v}건" for k, v in nz.items()) if nz else "없음") +
        " (D·E는 도입 집단 기준, 미도입 D·E 빈칸은 설계상 비해당).",
        "Y_total·W 완전사례 = 유효 표본 전원." if not any(miss[k] for k in cd.B_ITEMS + cd.C_ITEMS) else "")
    # 문항 텍스트 대조
    mism = [k for k in cd.ITEMS if xlsx_item_text(res.headers, k) != item_text.get(k)]
    add("검증", "—", "—",
        f"원본 머리글 문항 문구 ↔ 부록1 문항 전문 대조: 33문항 중 불일치 {len(mism)}건" + (f" ({', '.join(mism)})" if mism else "") + ".",
        "코드북 문항 전문은 부록1에서 읽음." if not mism else "부록1 기준으로 코드북 작성, 차이는 위 목록")
    t0, t1 = raw.timestamp.min(), raw.timestamp.max()
    vt = inc.merge(raw[["resp_id", "timestamp"]], on="resp_id").timestamp
    add("자료 특성(추가)", "—", "—",
        f"제출 기간: 전체 {t0:%Y-%m-%d %H:%M}∼{t1:%Y-%m-%d %H:%M}, 유효 응답 {vt.min():%Y-%m-%d %H:%M}∼{vt.max():%Y-%m-%d %H:%M}. "
        f"유효 표본 {len(inc)}명(도입 {len(ad)}, 미도입 {len(na)}) — 04장 모집 목표 60명과 다름.",
        "Google Forms 제출시각은 응답 소요시간으로 쓰지 않음(05장 제1절). 표본 수 서술 정합은 Task C에서.")
    return memos


def sheet_review(wb, memos):
    ws = wb.create_sheet("검토메모")
    header = ["번호", "구분", "resp_id", "원본 행", "내용", "처리", "상태"]
    write_table(ws, memos, header, freeze="A2")
    for r in range(2, len(memos) + 2):
        for c in range(1, 8):
            ws.cell(r, c).font = F_TEXT
            ws.cell(r, c).alignment = WRAP
    for c, w in zip("ABCDEFG", (6, 18, 22, 14, 70, 55, 18)):
        ws.column_dimensions[c].width = w
    return ws


def sheet_readme(ws, res: cd.BuildResult, sheets_info: list[tuple[str, str]], n_memos: int, generated: str):
    a = res.analysis
    inc = a[a.include == 1]
    lines = [
        ("설문 응답 정리본", None),
        ("목적", "Google Forms 응답 원본(도입/미도입 분기로 열이 두 묶음)을 한 열 세트로 합치고, 숫자 코딩·파생 척도·빈도/기술통계를 정리한 작업본."),
        ("원본 경로", str(cd.SRC_REL).replace("\\", "/")),
        ("원본 SHA256", res.sha256 + "  (읽기 전용 — 이 파일 생성 전후 불변 확인)"),
        ("생성 시각", generated),
        ("생성 스크립트", "tools/survey_analysis/build_workbook.py (코딩 규칙: tools/survey_analysis/coding.py)"),
        ("코딩 규칙 근거", "03.plan/261024_0850_설문응답_정리_분석_원고반영_계획.md 0.3·0.4절, 문항 전문 01.docs/부록1_설문지_양식.md, 변수 역할 01.docs/04_연구설계.md <표 4-3>"),
        ("응답 수", f"전체 {len(a)}건(도입 {int((a.group == 1).sum())}, 미도입 {int((a.group == 0).sum())}), "
                  f"분석 포함 include=1 {len(inc)}건(도입 {int((inc.group == 1).sum())}, 미도입 {int((inc.group == 0).sum())}). 시험 응답 2건 제외."),
        ("집단 기준", "group은 첫 질문의 자기보고 분기값이다. 원고 원칙상 분석용 집단은 조사 명부로 판정하므로 이 값은 '명부 대조 전 자기보고 기준'이다."),
        ("jamovi 투입", "분석자료 시트 또는 01.docs/설문결과/analysis/coded_data.csv(UTF-8 BOM)를 사용하고 include=1로 필터한다. 성명·현장명은 들어 있지 않다."),
        ("값 전용 시트", "분석자료는 설계상 값 전용(수식 없음)이다. 파생값은 Python으로 계산했으며 코딩자료의 수식과 같은 규칙이다(검증: 전 셀 일치)."),
        ("색 범례", "파란 글자 = 입력값(원문·코딩값·값 전용 자료), 검은 글자 = 수식, 노란 채움 = D함수 조건 범위(수정 금지)."),
        ("자료 특성 요약", "① 도입 집단 전원 남성(여성은 미도입에만) ② 소속–역할 중복(발주청≈공사감독)으로 설계행렬 순위 결함 가능 "
                     "③ 현장명·성명이 5건에만 있어 중복 확인은 타임스탬프·응답 벡터로만 가능(완전 동일 벡터 0건) "
                     "④ 집단은 명부 대조 전 자기보고 ⑤ 도입 집단 중 시스템 경로 작동(sys_path=1)은 소수 — 상세는 검토메모."),
        ("검토메모", f"{n_memos}개 항목. '사용자 확인 필요'·'확인 필요' 상태를 먼저 본다."),
    ]
    ws.cell(1, 1, lines[0][0]).font = F_TITLE
    r = 3
    for k, v in lines[1:]:
        ws.cell(r, 1, k).font = F_HEAD
        ws.cell(r, 2, v).font = F_TEXT
        ws.cell(r, 2).alignment = WRAP
        r += 1
    r += 1
    ws.cell(r, 1, "시트").font = F_HEAD
    ws.cell(r, 2, "설명").font = F_HEAD
    style_header_row(ws, r, 2)
    r += 1
    for name, desc in sheets_info:
        ws.cell(r, 1, name).font = F_TEXT
        ws.cell(r, 2, desc).font = F_TEXT
        ws.cell(r, 2).alignment = WRAP
        r += 1
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 120


SHEETS_INFO = [
    ("README", "이 안내"),
    ("원자료_통합", "응답 1건=1행(73행). 두 분기를 한 열 세트(A-1∼E-4)로 합친 원문 텍스트. 미도입의 D·E는 빈칸. 머리글 메모에 문항 전문."),
    ("코딩자료", "같은 73행의 숫자 코딩값(값) + 파생변수(수식: Y_total·Y1∼Y5·W·D_sub1∼4·D_total·E_total·sys_path). include·exclude_reason."),
    ("분석자료", "값 전용 jamovi/CSV 투입본. ASCII 변수명, 성명·현장명 제외. coded_data.csv와 동일."),
    ("코드북", "변수명·블록·문항 전문(부록1)·코딩 규칙·하위요인·분석 역할(<표 4-3>)·비고."),
    ("빈도_기술통계", "코딩자료를 참조하는 수식(COUNTIFS·D함수). 집단별(도입/미도입/전체, include=1) — 05장 <표 5-1>·<표 5-3> 골격 순서."),
    ("검토메모", "자료 특성·이상값·결측·중복 의심·제외·검증 결과. 각 건의 resp_id·원본 행·내용·처리·상태."),
]


# ---------------------------------------------------------------- 검증
def verify(out_xlsx: Path, res: cd.BuildResult, stats: StatsSheet, a: pd.DataFrame, out_csv: Path) -> dict:
    wbv = openpyxl.load_workbook(out_xlsx, data_only=True)
    wbf = openpyxl.load_workbook(out_xlsx)
    report: dict = {"sheets": {}, "errors": [], "forbidden": [], "formula_cells": 0}
    for ws in wbf.worksheets:
        report["sheets"][ws.title] = {"max_row": ws.max_row, "max_col": ws.max_column}
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    report["formula_cells"] += 1
                    up = c.value.upper()
                    if any(f in up for f in FORBIDDEN_FUNCS):
                        report["forbidden"].append(f"{ws.title}!{c.coordinate}")
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERROR_STRINGS:
                    report["errors"].append(f"{ws.title}!{c.coordinate}={c.value}")
    uncached = 0
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and wv[c.coordinate].value is None:
                    uncached += 1
    report["formula_cells_empty_result"] = uncached

    def same(x, y):
        if isblank(x) and isblank(y):
            return True
        if isblank(x) or isblank(y):
            return False
        return abs(float(x) - float(y)) <= 1e-9

    # 코딩자료 파생열 ↔ Python
    wc = wbv[CODED_SHEET]
    coded = res.coded
    deriv = {}
    for key in list(cd.DERIVED) + ["sys_path"]:
        col = CODED_COL[key]
        ok = nonempty = ok_inc = 0
        bad = []
        for i in range(N_ROWS):
            got = wc.cell(i + 2, col).value
            exp = coded.iloc[i][key]
            if same(got, exp):
                ok += 1
                if not isblank(exp):
                    nonempty += 1
                    if coded.iloc[i]["include"] == 1:
                        ok_inc += 1
            else:
                bad.append(f"{coded.iloc[i].resp_id}: excel={got} py={exp}")
        deriv[key] = {"match": ok, "of": N_ROWS, "nonempty_match": nonempty, "nonempty_match_include1": ok_inc,
                      "mismatch": bad}
    report["coded_derived_vs_python"] = deriv
    # 빈도_기술통계 ↔ Python
    ws = wbv["빈도_기술통계"]
    st_bad = [f"{k}: excel={ws[k].value} py={v}" for k, v in stats.expected.items() if not same(ws[k].value, v)]
    report["stats_vs_python"] = {"cells": len(stats.expected), "match": len(stats.expected) - len(st_bad), "mismatch": st_bad}
    # 분석자료 ↔ CSV ↔ Python
    csv = pd.read_csv(out_csv, encoding="utf-8-sig", dtype=str, keep_default_na=False)
    wa = wbv["분석자료"]
    cell_bad = 0
    for i in range(len(a)):
        for j, name in enumerate(a.columns):
            exp = a.iloc[i][name]
            got = wa.cell(i + 2, j + 1).value
            cv = csv.iloc[i][name]
            if name in ("resp_id", "A5_role_other", "A7_path_other", "exclude_reason"):
                ok = (isblank(exp) and isblank(got) and cv == "") or (str(exp) == str(got) == cv)
            else:
                ok = same(got, exp) and (same(float(cv), exp) if cv != "" else isblank(exp))
            cell_bad += 0 if ok else 1
    report["analysis_sheet_csv_python"] = {"cells": len(a) * len(a.columns), "mismatch": cell_bad,
                                           "csv_shape": list(csv.shape), "csv_columns_equal": list(csv.columns) == list(a.columns)}
    return report


# ---------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-recalc", action="store_true", help="Excel COM 재계산·검증을 건너뛴다")
    args = ap.parse_args(argv)

    sha_before = cd.sha256_of()
    if sha_before != cd.SRC_SHA256:
        raise SystemExit(f"원본 SHA256 불일치: {sha_before}")
    res = cd.build()
    item_text = appendix_items()
    missing = [k for k in cd.ITEMS if k not in item_text]
    if missing:
        raise SystemExit(f"부록1에서 문항 전문을 찾지 못함: {missing}")

    a = analysis_frame(res)
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    sheet_raw(wb, res, item_text)
    sheet_coded(wb, res)
    sheet_analysis(wb, a)
    sheet_codebook(wb, item_text)
    stats = sheet_stats(wb, res.coded)
    memos = review_items(res, item_text)
    sheet_review(wb, memos)
    generated = datetime.now(KST).strftime("%Y-%m-%d %H:%M (Asia/Seoul)")
    sheet_readme(readme, res, SHEETS_INFO, len(memos), generated)
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100

    out_xlsx = ROOT / OUT_XLSX_REL
    out_csv = ROOT / OUT_CSV_REL
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    a.to_csv(out_csv, index=False, encoding="utf-8-sig", lineterminator="\n")
    print(f"saved {OUT_XLSX_REL} / {OUT_CSV_REL}")

    report: dict = {"generated": generated, "sha256_before": sha_before}
    if not args.no_recalc:
        proc = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(RECALC_PS1),
                               "-Path", str(out_xlsx)], capture_output=True, text=True, encoding="utf-8", errors="replace")
        print(proc.stdout.strip(), proc.stderr.strip())
        if proc.returncode != 0:
            raise SystemExit("Excel COM 재계산 실패")
        report.update(verify(out_xlsx, res, stats, a, out_csv))
    sha_after = cd.sha256_of()
    report["sha256_after"] = sha_after
    report["sha256_unchanged"] = sha_before == sha_after == cd.SRC_SHA256
    inc = a[a.include == 1]
    report["counts"] = {"rows": len(a), "adopt": int((a.group == 1).sum()), "nonadopt": int((a.group == 0).sum()),
                        "include1": len(inc), "include1_adopt": int((inc.group == 1).sum()),
                        "include1_nonadopt": int((inc.group == 0).sum())}
    report["cross_fill_violations"] = len(res.cross_fill_violations)
    report["likert_failures"] = [list(map(str, x)) for x in res.likert_failures]
    report["category_failures"] = [list(map(str, x)) for x in res.category_failures]
    report["review_memo_items"] = len(memos)
    (ROOT / OUT_VERIFY_REL).write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    summary = {k: v for k, v in report.items() if k not in ("coded_derived_vs_python",)}
    print(json.dumps(summary, ensure_ascii=False, indent=1, default=str)[:4000])
    if "coded_derived_vs_python" in report:
        for k, v in report["coded_derived_vs_python"].items():
            print(f"  {k}: {v['match']}/{v['of']} 일치 (값 있는 셀 {v['nonempty_match']}, include=1 {v['nonempty_match_include1']}) 불일치 {len(v['mismatch'])}")
    ok = report["sha256_unchanged"] and (args.no_recalc or (
        not report["errors"] and not report["forbidden"] and not report["stats_vs_python"]["mismatch"]
        and all(not v["mismatch"] for v in report["coded_derived_vs_python"].values())
        and report["analysis_sheet_csv_python"]["mismatch"] == 0))
    print("VERIFY", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
