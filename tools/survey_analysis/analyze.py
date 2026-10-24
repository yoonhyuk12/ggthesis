"""설문 자료 통계 분석 — 05장 <표 5-1>∼<표 5-9>에 1:1 대응하는 결과 산출.

분석 설계의 단일 원본은 01.docs/04_연구설계.md 제5절이며, 이 스크립트는 그 절차를 그대로 옮긴다
(Type III ANCOVA, 기준범주 소속=시공사·성별=남성·역할=공사감독·총공사비=20억 원 미만,
조정 평균 = 공변량 완전사례 평균 + 범주형 동일 가중, 하위요인 Holm 보정, H2 W 평균중심화).

저장소 루트에서 실행 (1회 실행으로 전 산출물 재생성):
    PYTHONIOENCODING=utf-8 python tools/survey_analysis/analyze.py

산출: 01.docs/설문결과/analysis/{261024_분석결과.md, results.json, analyze_log.txt, figs/*.png}
"""
from __future__ import annotations

import io
import json
import math
import sys
import warnings
from datetime import datetime, timedelta, timezone
from itertools import product
from pathlib import Path

import numpy as np
import pandas as pd

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from tools.survey_analysis import coding as cd  # noqa: E402

import matplotlib  # noqa: E402
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import statsmodels.api as sm  # noqa: E402
import statsmodels.formula.api as smf  # noqa: E402
from scipy import stats  # noqa: E402
from statsmodels.stats.anova import anova_lm  # noqa: E402
from statsmodels.stats.diagnostic import het_breuschpagan  # noqa: E402
from statsmodels.stats.multitest import multipletests  # noqa: E402
from statsmodels.stats.outliers_influence import OLSInfluence, variance_inflation_factor  # noqa: E402

ROOT = cd.REPO_ROOT
OUT_DIR = ROOT / "01.docs/설문결과/analysis"
CSV_PATH = OUT_DIR / "coded_data.csv"
XLSX_PATH = ROOT / "01.docs/설문결과/261024_0850_설문응답_정리본.xlsx"
MD_PATH = OUT_DIR / "261024_분석결과.md"
JSON_PATH = OUT_DIR / "results.json"
LOG_PATH = OUT_DIR / "analyze_log.txt"
FIG_DIR = OUT_DIR / "figs"
KST = timezone(timedelta(hours=9))
ALPHA = 0.05

GROUP_LABEL = {1: "도입", 0: "미도입"}
# 범주형 통제요인과 기준범주(04장 제5절 제2항)
FACTORS = {
    "A1_affil": ("소속 구분", "A1", 1),
    "A2_sex": ("성별", "A2", 1),
    "A5_role": ("역할", "A5", 1),
    "A6_cost": ("총공사비 구간", "A6", 1),
}
COVARS = {"A3_age": "만 나이(1세 증가)", "A4_exp": "건설현장 경력(1년 증가)"}
SUBSCALES = {"Y1": "조기인지", "Y2": "대응 신속성", "Y3": "원격 파악", "Y4": "순찰부담 경감", "Y5": "위험 통제력"}
SCALE_LABEL = {"Y_total": "안전관리 실효성 전체", **SUBSCALES, "W": "안전관리 예산 제약성",
               "E_total": "경보 피로도 완화 인식", "D_total": "시스템 특성 인식 전체",
               "D_sub1": "인프라 범용성", "D_sub2": "탐지·제어 안정성", "D_sub3": "오탐 필터링",
               "D_sub4": "알림 전파 즉각성"}
RHS = ("D + A3_age + A4_exp + C(A1_affil, Treatment(1)) + C(A2_sex, Treatment(1))"
       " + C(A5_role, Treatment(1)) + C(A6_cost, Treatment(1))")
MODEL_VARS = ["D", "A3_age", "A4_exp", "A1_affil", "A2_sex", "A5_role", "A6_cost"]


# ---------------------------------------------------------------- 형식
def f2(x):
    return "—" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x:.2f}"


def f3(x):
    return "—" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x:.3f}"


def fp(p):
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return "—"
    return "<.001" if p < 0.001 else f"{p:.3f}".replace("0.", ".", 1)


def fpct(x):
    return "—" if x is None else f"{100 * x:.1f}"


def fci(lo, hi):
    return f"[{lo:.2f}, {hi:.2f}]"


def fnum(x):
    if x is None:
        return "—"
    return str(int(x)) if float(x).is_integer() else f"{x:.2f}"


def md_table(header, rows, align=None):
    align = align or ["---"] * len(header)
    out = ["| " + " | ".join(header) + " |", "|" + "|".join(align) + "|"]
    out += ["| " + " | ".join("" if v is None else str(v) for v in r) + " |" for r in rows]
    return "\n".join(out)


def jnum(x):
    """JSON용: numpy → python, NaN → None."""
    if isinstance(x, dict):
        return {str(k): jnum(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [jnum(v) for v in x]
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.floating, float)):
        return None if math.isnan(float(x)) else float(x)
    if isinstance(x, np.bool_):
        return bool(x)
    return x


# ---------------------------------------------------------------- 기초 통계 함수
def skew_kurt(x: np.ndarray):
    """표본 왜도 G1·초과첨도 G2(편의 보정, SPSS·jamovi와 같은 식)와 표준오차."""
    n = len(x)
    if n < 4:
        return dict(skew=None, kurt=None, se_skew=None, se_kurt=None)
    g1 = stats.skew(x, bias=False)
    g2 = stats.kurtosis(x, fisher=True, bias=False)
    se_s = math.sqrt(6 * n * (n - 1) / ((n - 2) * (n + 1) * (n + 3)))
    se_k = 2 * se_s * math.sqrt((n * n - 1) / ((n - 3) * (n + 5)))
    return dict(skew=float(g1), kurt=float(g2), se_skew=se_s, se_kurt=se_k)


def describe(x) -> dict:
    x = np.asarray(pd.Series(x).dropna(), dtype=float)
    n = len(x)
    d = dict(n=n, mean=float(x.mean()) if n else None, sd=float(x.std(ddof=1)) if n > 1 else None,
             median=float(np.median(x)) if n else None, min=float(x.min()) if n else None,
             max=float(x.max()) if n else None)
    d.update(skew_kurt(x))
    return d


def cronbach_alpha(df: pd.DataFrame):
    d = df.dropna()
    k = d.shape[1]
    if k < 2 or len(d) < 3:
        return None, len(d)
    item_var = d.var(ddof=1).sum()
    tot_var = d.sum(axis=1).var(ddof=1)
    return float(k / (k - 1) * (1 - item_var / tot_var)), len(d)


def corrected_item_total(df: pd.DataFrame):
    d = df.dropna()
    out = {}
    for c in d.columns:
        rest = d.drop(columns=c).sum(axis=1)
        out[c] = float(np.corrcoef(d[c], rest)[0, 1])
    return out


def alpha_if_deleted(df: pd.DataFrame):
    return {c: cronbach_alpha(df.drop(columns=c))[0] for c in df.columns}


def fisher_exact_2xk(table: np.ndarray):
    """2×k 교차표의 Fisher(Freeman–Halton) 정확검정 — 주변합 고정 전 표 열거.
    관측표 확률 이하인 표의 확률 합을 p로 한다(상대 허용오차 1e-7)."""
    table = np.asarray(table, dtype=int)
    keep = table.sum(axis=0) > 0
    t = table[:, keep]
    col = t.sum(axis=0)
    n1 = int(t[0].sum())
    N = int(col.sum())
    lfact = [math.lgamma(i + 1) for i in range(N + 1)]

    def logp(x):
        s = sum(lfact[c] - lfact[xi] - lfact[c - xi] for c, xi in zip(col, x))
        return s - (lfact[N] - lfact[n1] - lfact[N - n1])

    p_obs = logp(t[0])
    total, ntab = 0.0, 0
    k = len(col)
    suffix = np.cumsum(col[::-1])[::-1]

    def rec(j, remain, acc):
        nonlocal total, ntab
        if j == k - 1:
            if 0 <= remain <= col[j]:
                lp = logp(acc + [remain])
                ntab += 1
                if lp <= p_obs + 1e-7:
                    total += math.exp(lp)
            return
        rest_cap = suffix[j + 1]
        lo = max(0, remain - rest_cap)
        hi = min(col[j], remain)
        for v in range(lo, hi + 1):
            rec(j + 1, remain - v, acc + [v])

    rec(0, n1, [])
    return min(1.0, total), ntab, int(keep.sum())


def kmo_bartlett(df: pd.DataFrame):
    d = df.dropna()
    n, p = d.shape
    R = np.corrcoef(d.values, rowvar=False)
    inv = np.linalg.inv(R)
    partial = -inv / np.sqrt(np.outer(np.diag(inv), np.diag(inv)))
    off = ~np.eye(p, dtype=bool)
    r2, q2 = (R[off] ** 2).sum(), (partial[off] ** 2).sum()
    kmo = r2 / (r2 + q2)
    chi2 = -(n - 1 - (2 * p + 5) / 6) * math.log(np.linalg.det(R))
    dfb = p * (p - 1) // 2
    return dict(n=n, kmo=float(kmo), bartlett_chi2=float(chi2), bartlett_df=dfb,
                bartlett_p=float(stats.chi2.sf(chi2, dfb)), R=R)


def varimax(L, gamma=1.0, q=100, tol=1e-8):
    p, k = L.shape
    Rm = np.eye(k)
    d = 0
    for _ in range(q):
        d_old = d
        Lam = L @ Rm
        u, s, vh = np.linalg.svd(L.T @ (Lam ** 3 - (gamma / p) * Lam @ np.diag(np.diag(Lam.T @ Lam))))
        Rm = u @ vh
        d = s.sum()
        if d_old != 0 and d / d_old < 1 + tol:
            break
    return L @ Rm


def paf(R: np.ndarray, k: int, iters=1000, tol=1e-6):
    """반복 주축요인추출(PAF). 초기 공통성 = SMC."""
    h2 = 1 - 1 / np.diag(np.linalg.inv(R))
    for _ in range(iters):
        Rr = R.copy()
        np.fill_diagonal(Rr, h2)
        vals, vecs = np.linalg.eigh(Rr)
        idx = np.argsort(vals)[::-1][:k]
        L = vecs[:, idx] * np.sqrt(np.clip(vals[idx], 0, None))
        new = (L ** 2).sum(axis=1)
        if np.max(np.abs(new - h2)) < tol:
            h2 = new
            break
        h2 = new
    heywood = bool(np.any(h2 >= 1))
    return L, h2, heywood


def efa(df: pd.DataFrame, name: str):
    kb = kmo_bartlett(df)
    R = kb.pop("R")
    eig = np.sort(np.linalg.eigvalsh(R))[::-1]
    k = int((eig > 1).sum())
    L, h2, heywood = paf(R, k)
    if k > 1:
        L = varimax(L)
    L = L * np.sign(L.sum(axis=0))
    primary = np.abs(L).max(axis=1)
    ss = (L ** 2).sum(axis=0)
    out = dict(**kb, eigenvalues=eig.tolist(), n_factors=k, method="반복 주축요인추출(PAF), 초기 공통성 SMC"
               + (", varimax 회전" if k > 1 else ", 회전 없음(1요인)"),
               loadings={c: L[i].tolist() for i, c in enumerate(df.columns)},
               communalities={c: float(h2[i]) for i, c in enumerate(df.columns)},
               loading_min=float(primary.min()), loading_max=float(primary.max()),
               cum_var_pct=float(100 * ss.sum() / R.shape[0]), heywood=heywood, removed_items="없음")
    # 스크리 도표
    fig, ax = plt.subplots(figsize=(5, 3.4))
    ax.plot(range(1, len(eig) + 1), eig, "o-")
    ax.axhline(1, ls="--", c="gray", lw=0.8)
    ax.set_xlabel("요인 번호")
    ax.set_ylabel("고유값")
    ax.set_title(f"스크리 도표 — {name}")
    fig.tight_layout()
    fig.savefig(FIG_DIR / f"scree_{name}.png", dpi=150)
    plt.close(fig)
    return out


# ---------------------------------------------------------------- 모형
def fit_model(data: pd.DataFrame, dv: str, extra: str = ""):
    formula = f"{dv} ~ {RHS}{extra}"
    m = smf.ols(formula, data=data).fit()
    X = m.model.exog
    rank = int(np.linalg.matrix_rank(X))
    return m, formula, rank, X.shape[1]


def coef_table(m) -> list[dict]:
    ci = m.conf_int(alpha=ALPHA)
    rows = []
    for name in m.params.index:
        rows.append(dict(term=name, B=float(m.params[name]), SE=float(m.bse[name]),
                         ci_lo=float(ci.loc[name, 0]), ci_hi=float(ci.loc[name, 1]),
                         t=float(m.tvalues[name]), df=int(m.df_resid), p=float(m.pvalues[name])))
    return rows


TERM_KO = {
    "Intercept": "절편",
    "D": "도입(1)−미도입(0)",
    "A3_age": COVARS["A3_age"],
    "A4_exp": COVARS["A4_exp"],
    "W_c": "예산 제약성 W_c",
    "D:W_c": "D×W_c",
}


def term_label(term: str) -> str:
    if term in TERM_KO:
        return TERM_KO[term]
    for var, (lab, key, ref) in FACTORS.items():
        prefix = f"C({var}, Treatment({ref}))[T."
        if term.startswith(prefix):
            lvl = int(float(term[len(prefix):-1]))
            return f"{lab}: {cd.CAT_LABELS[key][lvl]} (기준 {cd.CAT_LABELS[key][ref]})"
    return term


def lvector(m, data: pd.DataFrame, d_value: float, weighting: str, w_c: float | None = None):
    """조정 평균의 선형결합 벡터. weighting='equal'(범주 동일 가중, jamovi EMM 기본) 또는 'proportional'(관측 비율)."""
    names = list(m.params.index)
    L = np.zeros(len(names))
    for i, nm in enumerate(names):
        if nm == "Intercept":
            L[i] = 1
        elif nm == "D":
            L[i] = d_value
        elif nm in ("A3_age", "A4_exp"):
            L[i] = data[nm].mean()
        elif nm == "W_c":
            L[i] = w_c
        elif nm == "D:W_c":
            L[i] = d_value * w_c
        else:
            for var, (lab, key, ref) in FACTORS.items():
                prefix = f"C({var}, Treatment({ref}))[T."
                if nm.startswith(prefix):
                    lvl = float(nm[len(prefix):-1])
                    if weighting == "equal":
                        L[i] = 1 / data[var].nunique()
                    else:
                        L[i] = float((data[var] == lvl).mean())
    return L


def lin_est(m, L):
    est = float(L @ m.params.values)
    se = float(math.sqrt(L @ m.cov_params().values @ L))
    tcrit = stats.t.ppf(1 - ALPHA / 2, m.df_resid)
    t = est / se if se > 0 else float("nan")
    return dict(est=est, se=se, ci_lo=est - tcrit * se, ci_hi=est + tcrit * se, t=t, df=int(m.df_resid),
                p=float(2 * stats.t.sf(abs(t), m.df_resid)))


def ancova(data: pd.DataFrame, dv: str, label: str):
    """H1 형 모형 1개 적합: Type III, 부분 η², 조정 평균(동일·비례 가중), 계수표, 식별 상태."""
    d = data.dropna(subset=[dv] + MODEL_VARS).copy()
    m, formula, rank, ncoef = fit_model(d, dv)
    a3 = anova_lm(m, typ=3)
    ss_g, ss_r = float(a3.loc["D", "sum_sq"]), float(a3.loc["Residual", "sum_sq"])
    res = dict(label=label, dv=dv, formula=formula, n=int(len(d)),
               n_adopt=int((d.D == 1).sum()), n_non=int((d.D == 0).sum()),
               n_coef=ncoef, rank=rank, df_resid=int(m.df_resid), full_rank=rank == ncoef,
               cond_number=float(np.linalg.cond(m.model.exog)),
               raw={GROUP_LABEL[g]: describe(d.loc[d.D == g, dv]) for g in (1, 0)},
               F=float(a3.loc["D", "F"]), df1=int(a3.loc["D", "df"]), df2=int(m.df_resid),
               p=float(a3.loc["D", "PR(>F)"]), partial_eta2=ss_g / (ss_g + ss_r),
               ss_group=ss_g, ss_resid=ss_r, r2=float(m.rsquared), adj_r2=float(m.rsquared_adj),
               coefs=coef_table(m),
               anova_type3={idx: dict(SS=float(r.sum_sq), df=float(r.df), F=None if math.isnan(r.F) else float(r.F),
                                      p=None if math.isnan(r["PR(>F)"]) else float(r["PR(>F)"]))
                            for idx, r in a3.iterrows()},
               covariate_ref={"A3_age": float(d.A3_age.mean()), "A4_exp": float(d.A4_exp.mean())})
    emm = {}
    for w in ("equal", "proportional"):
        e1 = lin_est(m, lvector(m, d, 1, w))
        e0 = lin_est(m, lvector(m, d, 0, w))
        diff = lin_est(m, lvector(m, d, 1, w) - lvector(m, d, 0, w))
        emm[w] = {"도입": e1, "미도입": e0, "차이": diff}
    res["emm"] = emm
    bD = next(c for c in res["coefs"] if c["term"] == "D")
    res["diff"] = bD
    res["check_emm_diff_equals_B"] = abs(emm["equal"]["차이"]["est"] - bD["B"]) < 1e-10
    res["check_F_equals_t2"] = abs(res["F"] - bD["t"] ** 2) < 1e-8
    return res, m, d


def diagnostics(m, d: pd.DataFrame, dv: str, tag: str, plots: bool = False):
    resid = m.resid.values
    fitted = m.fittedvalues.values
    n, p = m.model.exog.shape
    infl = OLSInfluence(m)
    cooks = infl.cooks_distance[0]
    lev = infl.hat_matrix_diag
    rstud = infl.resid_studentized_external
    ids = d.resp_id.values
    sw = stats.shapiro(resid)
    lev_g = [resid[(d.D == g).values] for g in (1, 0)]
    lev_mean = stats.levene(*lev_g, center="mean")
    lev_med = stats.levene(*lev_g, center="median")
    bp = het_breuschpagan(resid, m.model.exog)
    out = dict(
        tag=tag, n=n, n_params=p,
        shapiro_W=float(sw.statistic), shapiro_p=float(sw.pvalue),
        resid_skew=float(stats.skew(resid, bias=False)), resid_kurt=float(stats.kurtosis(resid, bias=False)),
        levene_mean_F=float(lev_mean.statistic), levene_mean_p=float(lev_mean.pvalue),
        levene_median_F=float(lev_med.statistic), levene_median_p=float(lev_med.pvalue),
        resid_sd_by_group={GROUP_LABEL[g]: float(r.std(ddof=1)) for g, r in zip((1, 0), lev_g)},
        bp_LM=float(bp[0]), bp_p=float(bp[1]),
        cooks_cut=4 / n, cooks_flag=[(str(i), float(c)) for i, c in zip(ids, cooks) if c > 4 / n],
        cooks_max=float(cooks.max()), cooks_max_id=str(ids[int(np.argmax(cooks))]),
        lev_cut=2 * p / n, lev_flag=[(str(i), float(h)) for i, h in zip(ids, lev) if h > 2 * p / n],
        rstud_flag=[(str(i), float(r)) for i, r in zip(ids, rstud) if abs(r) > 3],
        lev_one=[str(i) for i, h in zip(ids, lev) if h > 0.999],
    )
    if plots:
        fig, ax = plt.subplots(figsize=(4.5, 4.5))
        sm.qqplot(resid, line="s", ax=ax)
        ax.set_title(f"잔차 Q-Q 도표 — {tag}")
        fig.tight_layout()
        fig.savefig(FIG_DIR / f"qq_{dv}.png", dpi=150)
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(5.5, 4))
        for g, mk, col in ((1, "o", "tab:blue"), (0, "s", "tab:orange")):
            sel = (d.D == g).values
            ax.scatter(fitted[sel], resid[sel], marker=mk, c=col, alpha=0.7, label=GROUP_LABEL[g])
        ax.axhline(0, c="gray", lw=0.8)
        ax.set_xlabel("적합값")
        ax.set_ylabel("잔차")
        ax.set_title(f"적합값–잔차 — {tag}")
        ax.legend()
        fig.tight_layout()
        fig.savefig(FIG_DIR / f"resid_fitted_{dv}.png", dpi=150)
        plt.close(fig)
    return out


def group_coef(m, d: pd.DataFrame, base_ids, test: str = "t"):
    """집단 계수 D의 B·SE·95% CI·p와 기본 표본 대비 제외 사례."""
    ci = m.conf_int(alpha=ALPHA)
    return dict(n=int(len(d)), n_adopt=int((d.D == 1).sum()), n_non=int((d.D == 0).sum()),
                B=float(m.params["D"]), SE=float(m.bse["D"]), ci_lo=float(ci.loc["D", 0]), ci_hi=float(ci.loc["D", 1]),
                p=float(m.pvalues["D"]), df_resid=int(m.df_resid), test=test,
                excluded_ids=sorted(set(base_ids) - set(d.resp_id)))


def h1_subset_fit(d: pd.DataFrame):
    """H1 모형을 부분 표본에 다시 적합. 기준범주가 빠진 요인은 남은 최소 수준을 기준으로, 1수준만 남은 요인은 뺀다
    (집단 계수 B는 기준범주 선택과 무관)."""
    terms = ["D", "A3_age", "A4_exp"]
    for var, (_, _, ref) in FACTORS.items():
        lv = sorted(d[var].unique())
        if len(lv) > 1:
            terms.append(f"C({var}, Treatment({ref if ref in lv else int(lv[0])}))")
    m = smf.ols("Y_total ~ " + " + ".join(terms), data=d).fit()
    return m


def h1_robustness(d1: pd.DataFrame, m1, diag_h1: dict):
    """<표 5-6> 패널 E — H1 집단 차이의 사후 강건성 점검(결과 확인 후 수행, 모형·사례 선택에 쓰지 않음)."""
    ids = d1.resp_id.tolist()
    rob = {"note": "사후 점검. 결과 확인 뒤 수행했으며 사례·변수 선택에 사용하지 않는다. 기준범주가 빠진 부분 표본은 남은 최소 수준을 기준으로 재적합(B 불변)."}
    rob["base"] = group_coef(m1, d1, ids)
    m_hc3 = smf.ols(m1.model.formula, data=d1).fit(cov_type="HC3", use_t=True)
    rob["hc3"] = group_coef(m_hc3, d1, ids, test="t(HC3 강건 표준오차, 잔차 df)")
    rob["hc3"]["p_normal"] = float(2 * stats.norm.sf(abs(rob["hc3"]["B"] / rob["hc3"]["SE"])))  # statsmodels 기본(정규 근사) 참고값
    cooks_ids = [i for i, _ in diag_h1["cooks_flag"]]
    subsets = {
        "cooks_excluded": (~d1.resp_id.isin(cooks_ids), f"Cook's D>4/n({diag_h1['cooks_cut']:.4f}) 사례 제외"),
        "straightlining_excluded": (d1[cd.B_ITEMS + cd.C_ITEMS].nunique(axis=1) > 1,
                                    "일렬응답(B1∼B8·C1∼C4 12문항이 모두 같은 값) 제외"),
        "male_only": (d1.A2_sex == 1, "여성 제외(남성만)"),
        "exclude_cost1": (d1.A6_cost != 1, "총공사비 20억 원 미만 제외"),
        "exclude_role_other": (d1.A5_role != 4, "역할 기타 제외"),
    }
    for key, (mask, desc) in subsets.items():
        ds = d1[mask].copy()
        rob[key] = dict(**group_coef(h1_subset_fit(ds), ds, ids), condition=desc)
    rob["cooks_excluded"]["matches_table_5_5_flags"] = rob["cooks_excluded"]["excluded_ids"] == sorted(cooks_ids)
    m_un = smf.ols("Y_total ~ D", data=d1).fit()
    mw = stats.mannwhitneyu(d1.loc[d1.D == 1, "Y_total"], d1.loc[d1.D == 0, "Y_total"], alternative="two-sided")
    rob["unadjusted"] = dict(**group_coef(m_un, d1, ids), condition="공변량 없는 단순 비교(Y_total ~ D)",
                             mann_whitney_U=float(mw.statistic), mann_whitney_p=float(mw.pvalue))
    rob["mde80"] = dict(value=float((1.96 + 0.84) * rob["base"]["SE"]), multiplier=2.80, se_base=rob["base"]["SE"],
                        formula="MDE(80% 검정력, 양측 α=.05) = (1.96 + 0.84) × 기본 모형 SE(group)")
    return rob


def scatter_by_group(d, x, y, xlabel, ylabel, fname, title):
    fig, ax = plt.subplots(figsize=(5.5, 4))
    for g, mk, col in ((1, "o", "tab:blue"), (0, "s", "tab:orange")):
        s = d[d.D == g]
        ax.scatter(s[x], s[y], marker=mk, c=col, alpha=0.7, label=GROUP_LABEL[g])
        if len(s) > 2:
            b = np.polyfit(s[x], s[y], 1)
            xs = np.linspace(s[x].min(), s[x].max(), 20)
            ax.plot(xs, np.polyval(b, xs), c=col)
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.legend()
    fig.tight_layout()
    fig.savefig(FIG_DIR / fname, dpi=150)
    plt.close(fig)


def yang_dalton(p1: np.ndarray, p0: np.ndarray):
    """다범주 표준화 차이(Yang & Dalton, 2012): 기준범주를 뺀 K−1 비율 벡터의 마할라노비스 거리.
    공분산 S = (S1+S0)/2, Sg = diag(p) − p pᵀ."""
    a, b = p1[1:], p0[1:]
    S = ((np.diag(a) - np.outer(a, a)) + (np.diag(b) - np.outer(b, b))) / 2
    diff = a - b
    try:
        return float(math.sqrt(diff @ np.linalg.solve(S, diff)))
    except np.linalg.LinAlgError:
        return float("nan")


# ---------------------------------------------------------------- 메인 분석
def main():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "Malgun Gothic"
    plt.rcParams["axes.unicode_minus"] = False
    log = io.StringIO()
    R: dict = {"meta": {}}
    md: list[str] = []
    gaps: dict[str, list[str]] = {}

    with warnings.catch_warnings(record=True) as wlist:
        warnings.simplefilter("always")

        # ---------- 자료
        sha = cd.sha256_of()
        df = pd.read_csv(CSV_PATH, encoding="utf-8-sig")
        ref = cd.build().analysis
        same = (len(df) == len(ref)) and all(
            np.allclose(pd.to_numeric(df[c], errors="coerce").fillna(-999).values,
                        pd.to_numeric(ref[c], errors="coerce").fillna(-999).values)
            for c in df.columns if c not in ("resp_id", "A5_role_other", "A7_path_other", "exclude_reason"))
        inc = df[df.include == 1].copy()
        inc["D"] = inc["group"].astype(int)
        R["meta"] = dict(
            generated=datetime.now(KST).strftime("%Y-%m-%d %H:%M (Asia/Seoul)"),
            source_csv=str(CSV_PATH.relative_to(ROOT)).replace("\\", "/"),
            source_xlsx_sha256=sha, source_sha_ok=sha == cd.SRC_SHA256,
            csv_matches_coding_py=bool(same),
            n_rows=len(df), n_include=len(inc), n_include_adopt=int((inc.D == 1).sum()),
            n_include_non=int((inc.D == 0).sum()),
            software=f"Python {sys.version.split()[0]}, pandas {pd.__version__}, numpy {np.__version__}, "
                     f"scipy {__import__('scipy').__version__}, statsmodels {sm.__version__}",
            group_basis="자기보고 분기값(명부 대조 전)",
            alpha=ALPHA)
        G = {"도입": inc[inc.D == 1], "미도입": inc[inc.D == 0], "전체": inc}

        md.append("# 설문 자료 통계 분석 결과 (05장 <표 5-1>∼<표 5-9> 대응)\n")
        md.append(f"- 생성: {R['meta']['generated']} · 스크립트 `tools/survey_analysis/analyze.py` · 기계 판독 사본 `results.json`")
        md.append(f"- 자료: `{R['meta']['source_csv']}`(Task A 산출, coding.py 재생성본과 일치={same}), 원본 xlsx SHA256 `{sha[:12]}…`(불변={sha == cd.SRC_SHA256})")
        md.append(f"- 분석 대상: include=1 **{len(inc)}명**(도입 {len(G['도입'])}, 미도입 {len(G['미도입'])}). 시험 응답 2건 제외.")
        md.append("- **집단은 자기보고 분기값(조사 명부 대조 전)이다.** 통계 산출은 Python(statsmodels)이며 원고의 jamovi 재실행 대조가 남아 있다 `[확정 필요: jamovi 재실행으로 수치 대조]`.")
        md.append(f"- 소프트웨어: {R['meta']['software']}")
        md.append("- 표기: 비율 소수 1자리(%), M·SD·B·SE·CI 2자리, p 3자리(<.001), η²·α 3자리. 표 번호·패널은 05장과 같다.\n")

        # ================================================= 표 5-1
        t51 = {"panel_A": {}, "panel_B": {}, "panel_C": {}}
        rowsA = []
        for var, (lab, key, _) in FACTORS.items():
            rec = {}
            for code, clab in cd.CAT_LABELS[key].items():
                cell = {}
                row = [lab if code == 1 else "", clab]
                for g in ("도입", "미도입", "전체"):
                    s = G[g][var]
                    n, valid = int((s == code).sum()), int(s.notna().sum())
                    cell[g] = dict(n=n, pct=n / valid if valid else None)
                    row += [n, fpct(n / valid if valid else None)]
                rec[clab] = cell
                rowsA.append(row)
            miss = {g: dict(valid=int(G[g][var].notna().sum()), missing=int(G[g][var].isna().sum())) for g in G}
            rec["_denominator"] = miss
            rowsA.append(["", "(유효 응답 / 결측)"] + sum(([f"{miss[g]['valid']}", f"결측 {miss[g]['missing']}"] for g in G), []))
            t51["panel_A"][var] = rec
        rowsA.append(["**합계**", "", f"**{len(G['도입'])}**", "**100.0**", f"**{len(G['미도입'])}**", "**100.0**",
                      f"**{len(inc)}**", "**100.0**"])
        md.append("---\n## <표 5-1> 조사 대상의 인구통계학적 특성과 현장 운영 특성(집단별)\n")
        md.append("### 패널 A — 범주형 특성의 빈도와 비율\n")
        md.append(md_table(["구분", "항목", "도입 빈도(명)", "도입 비율(%)", "미도입 빈도(명)", "미도입 비율(%)", "전체 빈도(명)", "전체 비율(%)"],
                           rowsA, [":---:", ":---", "---:", "---:", "---:", "---:", "---:", "---:"]))
        md.append("\n주: 비율 분모는 항목별 유효 응답. 역할 결측 1건(R030, 미도입, A-5 '.' 응답)은 역할 비율 분모에서 빠진다.\n")
        rowsB = []
        for var, lab, unit in (("A3_age", "만 나이", "세"), ("A4_exp", "건설현장 경력", "년")):
            t51["panel_B"][var] = {}
            for g in ("도입", "미도입", "전체"):
                d_ = describe(G[g][var])
                t51["panel_B"][var][g] = d_
                rowsB.append([lab if g == "도입" else "", unit, g, d_["n"], f2(d_["mean"]), f2(d_["sd"]),
                              fnum(d_["median"]), fnum(d_["min"]), fnum(d_["max"])])
        md.append("### 패널 B — 만 나이·건설현장 경력\n")
        md.append(md_table(["연속 특성", "단위", "집단", "n", "평균(M)", "표준편차(SD)", "중위값", "최솟값", "최댓값"], rowsB,
                           [":---", ":---:", ":---:", "---:", "---:", "---:", "---:", "---:", "---:"]))
        md.append("\n주: 두 변수 모두 결측 0건.\n")
        rowsC = []
        fisher = {}
        for var, key, lab in (("A7_path", "A7", "A-7 담당자 부재 시 위험 인지 경로"),
                              ("A8_when", "A8", "A-8 순찰 간 위험 인지 시점"),
                              ("A9_lag", "A9", "A-9 인지 후 지시 전달 소요시간")):
            t51["panel_C"][var] = {}
            codes = list(cd.CAT_LABELS[key])
            for code in codes:
                cell = {}
                row = [lab if code == 1 else "", f"{'①②③④⑤'[code - 1]} {cd.CAT_LABELS[key][code]}"]
                for g in ("도입", "미도입", "전체"):
                    s = G[g][var]
                    n, valid = int((s == code).sum()), int(s.notna().sum())
                    cell[g] = dict(n=n, pct=n / valid)
                    row.append(f"{n} ({fpct(n / valid)})")
                t51["panel_C"][var][cd.CAT_LABELS[key][code]] = cell
                rowsC.append(row)
            tab = np.array([[int((G[g][var] == c).sum()) for c in codes] for g in ("도입", "미도입")])
            p_ex, ntab, ncols = fisher_exact_2xk(tab)
            chi = stats.chi2_contingency(tab[:, tab.sum(axis=0) > 0], correction=False)
            fisher[var] = dict(table=tab.tolist(), p_exact=p_ex, n_tables_enumerated=ntab, n_cols_used=ncols,
                               chi2=float(chi.statistic), chi2_df=int(chi.dof), chi2_p=float(chi.pvalue),
                               min_expected=float(chi.expected_freq.min()),
                               method="Fisher–Freeman–Halton 정확검정(2×k, 주변합 고정 전 표 열거)")
        t51["fisher"] = fisher
        ad = G["도입"]
        sp = dict(sys_path_1=int((ad.sys_path == 1).sum()), sys_path_0=int((ad.sys_path == 0).sum()),
                  a7_other=int((ad.A7_path == 5).sum()),
                  a7_12_or_a8_not3=int((ad.A7_path.isin([1, 2]) | (ad.A8_when != 3)).sum()),
                  ids_sys_path_1=ad.loc[ad.sys_path == 1, "resp_id"].tolist(),
                  ids_a7_other=ad.loc[ad.A7_path == 5, "resp_id"].tolist())
        t51["sys_path"] = sp
        md.append("### 패널 C — 현장 운영 특성(A-7∼A-9)의 집단별 분포\n")
        md.append(md_table(["문항", "보기", "도입 n(%)", "미도입 n(%)", "전체 n(%)"], rowsC, [":---", ":---", "---:", "---:", "---:"]))
        md.append("\n주: 비율 분모는 문항별 유효 응답(결측 0건). 집단 간 분포 차이의 Fisher 정확검정 p(참고값): "
                  + ", ".join(f"{k.split('_')[0].replace('A', 'A-')} p={fp(v['p_exact'])}" for k, v in fisher.items())
                  + ". 산출 방법은 2×k 교차표 주변합을 고정한 전체 표 열거(Fisher–Freeman–Halton)이며, 관측 빈도 0인 보기 열은 없다. "
                  "참고로 χ² 근사 p: " + ", ".join(f"{k.split('_')[0]} {fp(v['chi2_p'])}(최소 기대빈도 {v['min_expected']:.2f})" for k, v in fisher.items()) + ".")
        md.append(f"\n도입 집단 시스템 경로 작동(A-7 ③·④ 그리고 A-8 ③) **{sp['sys_path_1']}건**, 경로 미작동으로 따로 집계 **{sp['sys_path_0']}건**"
                  f"(그중 A-7 ⑤ 기타 {sp['a7_other']}건 — 판정 기준 `[확정 필요]`라 미작동으로 집계). 작동 사례: {', '.join(sp['ids_sys_path_1'])}.\n")
        gaps["5-1"] = ["패널 A 주의 '정상 운용 확인'(3개월 요건)·알림 경험 기록: 명부·운용 기록 미확보로 채울 수 없음.",
                       "집단은 자기보고 분기값이며 명부 대조 전이므로 개소 수(현장 수)=응답자 수로 가정만 가능(현장 ID 없음).",
                       "본문 '링크 안내 인원': 자료 없음. 제출 73·제외 2·분석 71은 채울 수 있음.",
                       "A-7 ⑤(기타) 판정 기준이 [확정 필요] 상태라 경로 미작동 집계는 잠정값."]
        R["table_5_1"] = t51

        # ================================================= 표 5-2
        t52 = {"panel_A": {}, "panel_B": {}}
        rows2 = []
        scales = [("Y_total", cd.B_ITEMS), ("Y1", ["B1"]), ("Y2", ["B2"]), ("Y3", ["B3", "B4"]),
                  ("Y4", ["B5", "B6"]), ("Y5", ["B7", "B8"]), ("W", cd.C_ITEMS)]
        for sc, items in scales:
            rec = {}
            ns, alphas, notes = [], [], []
            for g in ("전체", "도입", "미도입"):
                sub = G[g][items].dropna()
                ns.append(str(len(sub)))
                if len(items) == 1:
                    rec[g] = dict(n=len(sub), alpha=None)
                    continue
                a, n = cronbach_alpha(sub)
                r_ = dict(n=n, alpha=a)
                if len(items) == 2:
                    r_["r"] = float(np.corrcoef(sub[items[0]], sub[items[1]])[0, 1])
                else:
                    cm = sub.corr().values
                    off = cm[~np.eye(len(items), dtype=bool)]
                    r_["inter_item_r_mean"] = float(off.mean())
                    r_["inter_item_r_min"] = float(off.min())
                    r_["inter_item_r_max"] = float(off.max())
                    r_["corrected_item_total"] = corrected_item_total(sub)
                    r_["alpha_if_deleted"] = alpha_if_deleted(sub)
                rec[g] = r_
                alphas.append(f3(a))
            t52["panel_A"][sc] = rec
            if len(items) == 1:
                al = "해당 없음(단일 문항)"
                vc = G["전체"][items[0]].value_counts().sort_index()
                note = "응답 분포(전체) " + ", ".join(f"{int(k)}점 {int(v)}" for k, v in vc.items())
            elif len(items) == 2:
                al = "/".join(alphas)
                note = "문항 간 r(전체/도입/미도입) " + "/".join(f2(rec[g]["r"]) for g in ("전체", "도입", "미도입"))
            else:
                al = "/".join(alphas)
                e = rec["전체"]
                cit = e["corrected_item_total"]
                aid = e["alpha_if_deleted"]
                note = (f"문항 간 r 평균 {f2(e['inter_item_r_mean'])}(범위 {f2(e['inter_item_r_min'])}∼{f2(e['inter_item_r_max'])}); "
                        f"수정 문항–총점 r {f2(min(cit.values()))}∼{f2(max(cit.values()))}; "
                        f"문항 제거 시 α {f3(min(aid.values()))}∼{f3(max(aid.values()))}(전체 기준)")
            rows2.append([SCALE_LABEL[sc], len(items), "/".join(ns), al, note])
        md.append("---\n## <표 5-2> 측정도구의 신뢰도 검증 결과\n")
        md.append("### 패널 A — 전체 및 집단별 내적 일관성\n")
        md.append(md_table(["측정변수", "문항 수", "분석 개인 수(전체/도입/미도입)", "Cronbach's α(전체/도입/미도입)", "문항 간 관련성·점검 결과"],
                           rows2, [":---", "---:", ":---", ":---", ":---"]))
        md.append("\n주: 단일 문항 두 영역은 α를 산출하지 않는다. 문항 제거 시 α는 점검 정보이며 문항을 삭제하지 않았다.\n")
        efaY = efa(G["전체"][cd.B_ITEMS], "Y_B1-B8")
        efaW = efa(G["전체"][cd.C_ITEMS], "W_C1-C4")
        t52["panel_B"] = {"Y": efaY, "W": efaW}
        rowsE = []
        for lab, e in (("안전관리 실효성 문항(B-1∼B-8)", efaY), ("안전관리 예산 제약성 문항(C-1∼C-4)", efaW)):
            rowsE.append([lab, f3(e["kmo"]), f2(e["bartlett_chi2"]), e["bartlett_df"], fp(e["bartlett_p"]), e["n_factors"],
                          f"{f2(e['loading_min'])}∼{f2(e['loading_max'])}", f"{e['cum_var_pct']:.1f}", e["removed_items"]])
        md.append("### 패널 B — 탐색적 요인분석(참고 산출, `[확정 필요]` 유지)\n")
        md.append(md_table(["분석 대상", "KMO", "Bartlett χ²", "df", "p", "추출 요인 수", "요인적재값 범위", "누적 설명분산(%)", "제거 문항"],
                           rowsE, [":---", "---:", "---:", "---:", "---:", "---:", ":---:", "---:", ":---:"]))
        md.append(f"\n주: 전체 유효 표본(n={efaY['n']}). 요인 수 = 상관행렬 고유값>1 개수(스크리 도표 `figs/scree_*.png` 병행), "
                  f"추출 = {efaY['method']}(Y) / {efaW['method']}(W). 적재값 범위는 문항별 최대 절대 적재값의 범위. "
                  f"고유값 Y: {', '.join(f'{v:.2f}' for v in efaY['eigenvalues'])}; W: {', '.join(f'{v:.2f}' for v in efaW['eigenvalues'])}. "
                  f"Heywood 사례 Y={efaY['heywood']}, W={efaW['heywood']}. 04장 제5절은 EFA 적용 여부·추출·회전·요인 수 기준을 "
                  "`[확정 필요]`로 두고 있으므로 이 패널은 기준 확정 전 참고값이다.\n")
        if efaY["n_factors"] > 1:
            md.append("Y 8문항 회전 적재값(varimax): " + "; ".join(
                f"{k} " + "/".join(f2(v) for v in vals) for k, vals in efaY["loadings"].items()) + "\n")
        gaps["5-2"] = ["본문 전문가 CVR·인지면접·예비조사·채점안 고정일: 자료 없음([DATA PENDING] 유지).",
                       "패널 B: EFA 포함 여부·추출·회전 기준이 04장에서 [확정 필요] — 값은 참고 산출이며 패널 유지/삭제 결정은 사용자 몫.",
                       "통계적 측정동일성(집단 간) 검증은 수행하지 않음(표본 규모상 설계에 없음)."]
        R["table_5_2"] = t52

        # ================================================= 표 5-3
        t53 = {"panel_A": {}, "panel_B": {}}
        rows3 = []
        for sc in ("Y_total", "Y1", "Y2", "Y3", "Y4", "Y5", "W", "E_total"):
            t53["panel_A"][sc] = {}
            groups = ("도입", "미도입", "전체") if sc != "E_total" else ("도입",)
            for g in groups:
                d_ = describe(G[g][sc])
                t53["panel_A"][sc][g] = d_
                rows3.append([SCALE_LABEL[sc] if g == groups[0] else "", g, d_["n"], f2(d_["mean"]), f2(d_["sd"]),
                              f2(d_["skew"]), f2(d_["kurt"])])
        md.append("---\n## <표 5-3> 주요 변수의 기술통계(집단별)\n")
        md.append("### 패널 A — 공통 척도와 E 완전응답 평균점수의 분포\n")
        md.append(md_table(["변수", "집단", "n", "평균(M)", "표준편차(SD)", "왜도", "첨도"], rows3,
                           [":---", ":---:", "---:", "---:", "---:", "---:", "---:"]))
        md.append("\n주: 왜도 = 편의 보정 표본 왜도 G1 = √(n(n−1))/(n−2)·m3/m2^1.5, 첨도 = 편의 보정 초과첨도 "
                  "G2 = (n−1)/((n−2)(n−3))·[(n+1)·m4/m2² − 3(n−1)] (정규분포 0, SPSS·jamovi와 같은 식). "
                  "05장 원고 표는 Y·하위요인·W의 '전체' 행을 Y 전체에만 두지만 여기서는 모든 공통 척도에 전체 행을 참고로 덧붙였다. "
                  "D·E 하위척도 기술통계는 results.json `table_5_3.reference_D`.\n")
        t53["reference_D"] = {sc: describe(G["도입"][sc]) for sc in ("D_total", "D_sub1", "D_sub2", "D_sub3", "D_sub4")}
        rows3b = []
        nad = len(G["도입"])
        for it in cd.D_ITEMS + cd.E_ITEMS:
            d_ = describe(G["도입"][it])
            t53["panel_B"][it] = dict(**d_, nonresponse=nad - d_["n"])
            rows3b.append([f"{it[0]}-{it[1:]}", d_["n"], "—", nad - d_["n"], f2(d_["mean"]), f2(d_["sd"])])
        md.append("### 패널 B — 도입 집단 전용 D·E 문항별 응답 상태와 기술통계\n")
        md.append(md_table(["문항", "유효 응답 수", "경험 부재 비해당 수", "일반 무응답 수", "평균(M)", "표준편차(SD)"], rows3b,
                           [":---", "---:", "---:", "---:", "---:", "---:"]))
        md.append(f"\n주: 전용부 응답 대상 도입 개인 수 {nad}명과 대조(유효+무응답 = {nad}).\n")
        diffs = {sc: t53["panel_A"][sc]["도입"]["mean"] - t53["panel_A"][sc]["미도입"]["mean"] for sc in SUBSCALES}
        big = max(diffs, key=lambda k: abs(diffs[k]))
        t53["raw_mean_diff_subscales"] = diffs
        t53["largest_raw_diff_subscale"] = big
        md.append(f"원평균 차이(도입−미도입)가 가장 큰 하위요인: **{SCALE_LABEL[big]}**({f2(diffs[big])}점). 하위요인별 원평균 차이: "
                  + ", ".join(f"{SCALE_LABEL[k]} {f2(v)}" for k, v in diffs.items()) + ".\n")
        gaps["5-3"] = ["없음(패널 A·B 모두 채움). 분포 판단의 문장화는 Task C 몫."]
        R["table_5_3"] = t53

        # ================================================= 표 5-4
        t54 = {"panel_A": {}, "panel_B": {}}
        rows4 = []
        for var, lab in (("A3_age", "만 나이(세)"), ("A4_exp", "건설현장 경력(년)")):
            a, b = G["도입"][var].dropna(), G["미도입"][var].dropna()
            sp_ = math.sqrt(((len(a) - 1) * a.var() + (len(b) - 1) * b.var()) / (len(a) + len(b) - 2))
            lo, hi = max(a.min(), b.min()), min(a.max(), b.max())
            out_a = int(((a < lo) | (a > hi)).sum())
            out_b = int(((b < lo) | (b > hi)).sum())
            rec = dict(adopt=describe(a), non=describe(b), mean_diff=float(a.mean() - b.mean()), pooled_sd=sp_,
                       smd=float((a.mean() - b.mean()) / sp_), common_lo=float(lo), common_hi=float(hi),
                       outside_common_adopt=out_a, outside_common_non=out_b,
                       smd_formula="(M도입−M미도입)/합동 SD, 합동 SD=√[((n1−1)s1²+(n0−1)s0²)/(n1+n0−2)]")
            t54["panel_A"][var] = rec
            rows4.append([lab, "개인", f"{len(a)}·{f2(a.mean())}({f2(a.std())})·{fnum(a.min())}∼{fnum(a.max())}",
                          f"{len(b)}·{f2(b.mean())}({f2(b.std())})·{fnum(b.min())}∼{fnum(b.max())}",
                          f2(rec["mean_diff"]), f2(rec["smd"]),
                          f"{fnum(lo)}∼{fnum(hi)} (범위 밖 도입 {out_a}·미도입 {out_b}명)"])
        md.append("---\n## <표 5-4> 두 집단의 사전 특성 분포·균형 진단 결과\n")
        md.append("### 패널 A — 연속 공변량의 분포\n")
        md.append(md_table(["사전 특성", "집계 단위", "도입 n·M(SD)·범위", "미도입 n·M(SD)·범위", "평균 차이", "표준화 차이", "공통 관측범위"],
                           rows4, [":---", ":---", ":---", ":---", ":---", ":---", ":---"]))
        md.append("\n주: 평균 차이 = 도입−미도입. 표준화 차이 = 평균 차이 / 합동 표준편차(n−1 가중). 공통 관측범위 = 두 집단 관측범위의 교집합.\n")
        rows4b = []
        for var, (lab, key, ref_) in FACTORS.items():
            codes = list(cd.CAT_LABELS[key])
            a, b = G["도입"][var].dropna(), G["미도입"][var].dropna()
            p1 = np.array([(a == c).mean() for c in codes])
            p0 = np.array([(b == c).mean() for c in codes])
            cat_smd = {}
            for i, c in enumerate(codes):
                den = math.sqrt((p1[i] * (1 - p1[i]) + p0[i] * (1 - p0[i])) / 2)
                cat_smd[cd.CAT_LABELS[key][c]] = float((p1[i] - p0[i]) / den) if den > 0 else float("nan")
            yd = yang_dalton(p1, p0)
            sparse = [f"{GROUP_LABEL[g]} {cd.CAT_LABELS[key][c]} {int((G[GROUP_LABEL[g]][var] == c).sum())}"
                      for g in (1, 0) for c in codes if (G[GROUP_LABEL[g]][var] == c).sum() < 5]
            empty = [f"{GROUP_LABEL[g]} {cd.CAT_LABELS[key][c]}" for g in (1, 0) for c in codes
                     if (G[GROUP_LABEL[g]][var] == c).sum() == 0]
            t54["panel_B"][var] = dict(n_adopt=len(a), n_non=len(b), p_adopt=p1.tolist(), p_non=p0.tolist(),
                                       prop_diff=(p1 - p0).tolist(), smd_yang_dalton=yd, smd_by_category=cat_smd,
                                       max_abs_cat_smd=float(np.nanmax(np.abs(list(cat_smd.values())))),
                                       sparse_lt5=sparse, empty_cells=empty)
            unit = "현장" if var == "A6_cost" else "개인"
            rows4b.append([lab, unit,
                           "; ".join(f"{cd.CAT_LABELS[key][c]} {int((a == c).sum())}({fpct(p1[i])})" for i, c in enumerate(codes)),
                           "; ".join(f"{cd.CAT_LABELS[key][c]} {int((b == c).sum())}({fpct(p0[i])})" for i, c in enumerate(codes)),
                           "; ".join(f"{cd.CAT_LABELS[key][c]} {100 * (p1[i] - p0[i]):+.1f}%p" for i, c in enumerate(codes)),
                           f"{f2(yd)} (범주별 최대 절대 d {f2(t54['panel_B'][var]['max_abs_cat_smd'])})",
                           ("빈 칸: " + ", ".join(empty) + "; " if empty else "") + ("n<5: " + ", ".join(sparse) if sparse else "희소범주 없음")])
        md.append("### 패널 B — 범주형 통제요인의 분포\n")
        md.append(md_table(["사전 특성", "집계 단위", "도입 범주별 n(%)", "미도입 범주별 n(%)", "범주별 비율 차이", "표준화 차이", "공통 관측범위·희소범주"],
                           rows4b, [":---", ":---", ":---", ":---", ":---", ":---", ":---"]))
        md.append("\n주: 표준화 차이 = Yang & Dalton(2012)의 다범주 표준화 차이(기준범주를 뺀 K−1 비율 벡터의 마할라노비스 거리, "
                  "공분산은 두 집단 다항 공분산의 평균; 2범주에서는 |p1−p0|/√[(p1q1+p0q0)/2]와 같다)이며 부호가 없다. "
                  "괄호 안은 범주별 이분 표준화 차이의 최대 절댓값. 분모는 변수별 유효 응답(역할 미도입 38, 결측 1). "
                  "[CITE_TODO: Yang & Dalton(2012) 다범주 표준화 차이 — 참고문헌 목록 미등재]\n")
        r_ae = float(inc[["A3_age", "A4_exp"]].corr().iloc[0, 1])
        xt = pd.crosstab(inc.A1_affil.map(cd.CAT_LABELS["A1"]), inc.A5_role.map(cd.CAT_LABELS["A5"]))
        t54["age_exp_r"] = r_ae
        t54["affil_x_role"] = {str(i): {str(c): int(v) for c, v in r.items()} for i, r in xt.iterrows()}
        md.append(f"만 나이–경력 상관 r={f2(r_ae)}(include=1, n={len(inc)}). 소속×역할 교차표(include=1, 역할 결측 제외):\n")
        md.append(md_table(["소속 \\ 역할"] + list(xt.columns), [[i] + [int(v) for v in r] for i, r in xt.iterrows()]))
        md.append("")
        gaps["5-4"] = ["본문 현장 단위 인프라·현장 조건 분포(CCTV 보유·원격 열람·다른 스마트 장비·공종·인력), 도입 현장 운용 설정: 자료 없음.",
                       "'불균형·외삽 위험' 판단 문장화는 Task C 몫(수치는 위 표)."]
        R["table_5_4"] = t54

        # ================================================= 표 5-6 (H1) + 5-5 진단
        h1, m1, d1 = ancova(inc, "Y_total", "H1")
        diag_h1 = diagnostics(m1, d1, "Y_total", "H1 Y_total", plots=True)
        # 기울기 동질성
        m_slope = smf.ols(f"Y_total ~ {RHS} + D:A3_age + D:A4_exp", data=d1).fit()
        a3s = anova_lm(m_slope, typ=3)
        jt = m_slope.f_test("D:A3_age = 0, D:A4_exp = 0")
        slope = dict(D_x_age_F=float(a3s.loc["D:A3_age", "F"]), D_x_age_p=float(a3s.loc["D:A3_age", "PR(>F)"]),
                     D_x_exp_F=float(a3s.loc["D:A4_exp", "F"]), D_x_exp_p=float(a3s.loc["D:A4_exp", "PR(>F)"]),
                     joint_F=float(jt.fvalue), joint_df=(int(jt.df_num), int(jt.df_denom)), joint_p=float(jt.pvalue),
                     slopes_by_group={GROUP_LABEL[g]: {v: float(np.polyfit(d1.loc[d1.D == g, v], d1.loc[d1.D == g, "Y_total"], 1)[0])
                                                       for v in ("A3_age", "A4_exp")} for g in (1, 0)})
        # 선형성 진단: 제곱항 추가 F
        m_quad = smf.ols(f"Y_total ~ {RHS} + I(A3_age**2) + I(A4_exp**2)", data=d1).fit()
        qt = m_quad.f_test("I(A3_age ** 2) = 0, I(A4_exp ** 2) = 0")
        lin = dict(r_age_Y={g: float(np.corrcoef(G[g].A3_age, G[g].Y_total)[0, 1]) for g in G},
                   r_exp_Y={g: float(np.corrcoef(G[g].A4_exp, G[g].Y_total)[0, 1]) for g in G},
                   quad_F=float(qt.fvalue), quad_df=(int(qt.df_num), int(qt.df_denom)), quad_p=float(qt.pvalue))
        scatter_by_group(d1, "A3_age", "Y_total", "만 나이(세)", "Y_total", "scatter_age_Y_by_group.png", "만 나이–Y (집단별 회귀선)")
        scatter_by_group(d1, "A4_exp", "Y_total", "건설현장 경력(년)", "Y_total", "scatter_exp_Y_by_group.png", "경력–Y (집단별 회귀선)")
        # 공선성
        X = pd.DataFrame(m1.model.exog, columns=m1.model.exog_names)
        vif = {c: float(variance_inflation_factor(X.values, i)) for i, c in enumerate(X.columns) if c != "Intercept"}
        # 하위요인 단일문항 분포
        ceil_floor = {}
        for sc in SUBSCALES:
            ceil_floor[sc] = {g: dict(ceiling_pct=float((G[g][sc] == 5).mean()), floor_pct=float((G[g][sc] == 1).mean()))
                              for g in ("도입", "미도입")}
        # 민감도 분석
        sens_data = inc[(inc.D == 0) | (inc.sys_path == 1)]
        sens, m_sens, d_sens = ancova(sens_data, "Y_total", "H1 민감도(sys_path=0 도입 제외)")
        sens["excluded_adopt_n"] = int(((inc.D == 1) & (inc.sys_path == 0)).sum())
        sens["empty_levels_in_adopt"] = {var: sorted(set(d1[var].unique()) - set(d_sens.loc[d_sens.D == 1, var].unique()))
                                         for var in FACTORS}
        # 민감도 분석 보조: 시스템 경로 작동/미작동 도입 현장과 미도입의 Y 원평균(H1 완전사례)
        sens["path_split"] = {lab: dict(n=int(len(s)), mean=float(s.Y_total.mean()), sd=float(s.Y_total.std(ddof=1)))
                              for lab, s in (("adopt_sys_path_1", d1[(d1.D == 1) & (d1.sys_path == 1)]),
                                             ("adopt_sys_path_0", d1[(d1.D == 1) & (d1.sys_path == 0)]),
                                             ("non_adopt", d1[d1.D == 0]))}
        R["table_5_6"] = dict(H1=h1, sensitivity=sens, robustness=h1_robustness(d1, m1, diag_h1))

        # ================================================= 표 5-7
        subs = {}
        sub_models = {}
        for sc, lab in SUBSCALES.items():
            r_, m_, d_ = ancova(inc, sc, lab)
            r_["diag"] = diagnostics(m_, d_, sc, lab)
            subs[sc] = r_
            sub_models[sc] = m_
        raw_p = [subs[sc]["p"] for sc in SUBSCALES]
        holm = multipletests(raw_p, alpha=ALPHA, method="holm")[1]
        for sc, hp in zip(SUBSCALES, holm):
            subs[sc]["p_holm"] = float(hp)
        R["table_5_7"] = subs

        # ================================================= 표 5-8 (H2)
        d2 = inc.dropna(subset=["Y_total", "W"] + MODEL_VARS).copy()
        w_mean, w_sd = float(d2.W.mean()), float(d2.W.std(ddof=1))
        d2["W_c"] = d2.W - w_mean
        m2 = smf.ols(f"Y_total ~ {RHS} + W_c + D:W_c", data=d2).fit()
        rank2, ncoef2 = int(np.linalg.matrix_rank(m2.model.exog)), m2.model.exog.shape[1]
        a32 = anova_lm(m2, typ=3)
        w_rng = {GROUP_LABEL[g]: (float(d2.loc[d2.D == g, "W"].min()), float(d2.loc[d2.D == g, "W"].max())) for g in (1, 0)}
        c_lo, c_hi = max(v[0] for v in w_rng.values()), min(v[1] for v in w_rng.values())
        cond = {}
        for lab, w in (("평균−1SD", -w_sd), ("평균", 0.0), ("평균+1SD", w_sd)):
            wv = w_mean + w
            inside = c_lo <= wv <= c_hi
            L1 = lvector(m2, d2, 1, "equal", w)
            L0 = lvector(m2, d2, 0, "equal", w)
            rec = dict(W_c=w, W=wv, inside_common=inside)
            if inside:
                rec.update(adopt=lin_est(m2, L1), non=lin_est(m2, L0), diff=lin_est(m2, L1 - L0))
                rec["check_diff_equals_bD_plus_bDW_w"] = abs(rec["diff"]["est"] - (m2.params["D"] + m2.params["D:W_c"] * w)) < 1e-10
                # 비례 가중(관측 분포 기준) 조건부 조정 평균 — 기존 동일 가중 키는 그대로 두고 새 키로 추가
                P1 = lvector(m2, d2, 1, "proportional", w)
                P0 = lvector(m2, d2, 0, "proportional", w)
                rec.update(adopt_proportional=lin_est(m2, P1), non_proportional=lin_est(m2, P0),
                           diff_proportional=lin_est(m2, P1 - P0))
                rec["check_proportional_diff_equals_diff"] = abs(rec["diff_proportional"]["est"] - rec["diff"]["est"]) < 1e-10
            else:
                rec["excluded_reason"] = f"W={wv:.2f}가 공통 관측범위 {c_lo:.2f}∼{c_hi:.2f} 밖"
            cond[lab] = rec
        int_row = next(c for c in coef_table(m2) if c["term"] == "D:W_c")
        h2 = dict(n=int(len(d2)), n_adopt=int((d2.D == 1).sum()), n_non=int((d2.D == 0).sum()),
                  W_mean=w_mean, W_sd=w_sd, W_range_by_group=w_rng, W_common=(c_lo, c_hi),
                  n_coef=ncoef2, rank=rank2, df_resid=int(m2.df_resid), full_rank=rank2 == ncoef2,
                  coefs=coef_table(m2), interaction=int_row,
                  interaction_F=float(a32.loc["D:W_c", "F"]), interaction_df=(int(a32.loc["D:W_c", "df"]), int(m2.df_resid)),
                  interaction_p=float(a32.loc["D:W_c", "PR(>F)"]),
                  partial_eta2_interaction=float(a32.loc["D:W_c", "sum_sq"] / (a32.loc["D:W_c", "sum_sq"] + a32.loc["Residual", "sum_sq"])),
                  conditional=cond, r_W_Y={g: float(np.corrcoef(G[g].W, G[g].Y_total)[0, 1]) for g in G},
                  note_W_c_coef="W_c 계수는 D=0(미도입) 집단의 W 기울기(처리 부호화). 도입 집단 기울기 = W_c + D×W_c.")
        h2["slope_adopt"] = lin_est(m2, np.array([1.0 if n in ("W_c", "D:W_c") else 0.0 for n in m2.params.index]))
        h2["diag"] = diagnostics(m2, d2, "Y_total_H2", "H2")
        scatter_by_group(d2, "W", "Y_total", "예산 제약성 W", "Y_total", "scatter_W_Y_by_group.png", "W–Y (집단별 회귀선)")
        h2["mde80"] = dict(value=float((1.96 + 0.84) * int_row["SE"]), multiplier=2.80, se_base=int_row["SE"],
                           formula="MDE(80% 검정력, 양측 α=.05) = (1.96 + 0.84) × SE(D×W_c)")
        # 사후 점검 — H2 모형의 HC3 강건 표준오차(Breusch–Pagan 이분산 신호 대응). 기존 키는 그대로 두고 새 키로 추가
        m2_hc3 = smf.ols(m2.model.formula, data=d2).fit(cov_type="HC3", use_t=True)
        ci2_hc3 = m2_hc3.conf_int(alpha=ALPHA)
        h2["hc3"] = dict(note="사후 점검. 결과 확인 뒤 수행했으며 판정 기준은 기본 모형(<표 5-8> 패널 C)이다. p는 잔차 df의 t분포, p_normal은 정규 근사 참고값.",
                         df_resid=int(m2_hc3.df_resid), test="t(HC3 강건 표준오차, 잔차 df)",
                         terms={t: dict(B=float(m2_hc3.params[t]), SE=float(m2_hc3.bse[t]),
                                        ci_lo=float(ci2_hc3.loc[t, 0]), ci_hi=float(ci2_hc3.loc[t, 1]),
                                        p=float(m2_hc3.pvalues[t]),
                                        p_normal=float(2 * stats.norm.sf(abs(m2_hc3.params[t] / m2_hc3.bse[t]))))
                                for t in ("D", "W_c", "D:W_c")})
        h2["hc3"]["conditional_diff"] = {lab: lin_est(m2_hc3, lvector(m2, d2, 1, "equal", rec["W_c"]) - lvector(m2, d2, 0, "equal", rec["W_c"]))
                                         for lab, rec in cond.items() if rec["inside_common"]}
        R["table_5_8"] = h2

        # ================================================= A-9 판정
        f9 = fisher["A9_lag"]
        p9a = np.array(f9["table"][0]) / sum(f9["table"][0])
        p9n = np.array(f9["table"][1]) / sum(f9["table"][1])
        mw = stats.mannwhitneyu(G["도입"].A9_lag, G["미도입"].A9_lag, alternative="two-sided")
        a9 = dict(p_fisher=f9["p_exact"], prop_adopt=p9a.tolist(), prop_non=p9n.tolist(),
                  cum_adopt=np.cumsum(p9a).tolist(), cum_non=np.cumsum(p9n).tolist(),
                  mean_code_adopt=float(G["도입"].A9_lag.mean()), mean_code_non=float(G["미도입"].A9_lag.mean()),
                  exploratory_mann_whitney=dict(U=float(mw.statistic), p=float(mw.pvalue),
                                                note="탐색적 부가 분석(설계 외). 순서 척도 위치 비교 참고값"))
        a9_sig = f9["p_exact"] < ALPHA
        a9["provisional_judgement"] = ("있음" if a9_sig and a9["cum_adopt"][0] > a9["cum_non"][0] else
                                       "역방향 차이" if a9_sig else "없음(Fisher p≥.05, 잠정)")
        h1_sig = h1["p"] < ALPHA
        h1_judge = ("지지(차이 양수, p<.05)" if h1_sig and h1["diff"]["B"] > 0 else
                    "역방향 차이" if h1_sig else "기각(차이 미확인)")
        a9["provisional_combination"] = (f"A-9 {'있음' if a9_sig else '없음'} × H1 {'있음' if (h1_sig and h1['diff']['B'] > 0) else '없음'}")
        R["a9_judgement"] = a9

        # ================================================= 표 5-5 작성
        md.append("---\n## <표 5-5> 공분산분석 가정 점검 결과 (H1 모형 기준, 하위요인·H2 요약 병기)\n")
        dg = diag_h1
        sub_norm = "; ".join(f"{SUBSCALES[s]} W={subs[s]['diag']['shapiro_W']:.3f}, p={fp(subs[s]['diag']['shapiro_p'])}" for s in SUBSCALES)
        sub_bp = "; ".join(f"{SUBSCALES[s]} p={fp(subs[s]['diag']['bp_p'])}" for s in SUBSCALES)
        cf = "; ".join(f"{SUBSCALES[s]} 도입 {fpct(ceil_floor[s]['도입']['ceiling_pct'])}%/{fpct(ceil_floor[s]['도입']['floor_pct'])}%, "
                       f"미도입 {fpct(ceil_floor[s]['미도입']['ceiling_pct'])}%/{fpct(ceil_floor[s]['미도입']['floor_pct'])}%" for s in SUBSCALES)
        vif_max = max(vif, key=vif.get)
        rows5 = [
            ["관측의 독립성", "현장당 1명 여부·현장 ID·담당자 중복",
             "현장 ID·명부 미확보로 현장 중복은 확인 불가. 전 문항 응답 벡터 완전 동일 쌍 0건(Task A 검토메모). 2분 이내 근접 제출 24쌍은 응답이 서로 달라 중복 근거 없음.",
             "자료만으로 독립성 확인 불가 — 명부 대조 후 확정. 잠정적으로 독립 관측으로 분석."],
            ["선형성", "나이·경력–Y 상관, 제곱항 추가 F, 적합값–잔차 도표",
             f"r(나이,Y) 도입 {f2(lin['r_age_Y']['도입'])}·미도입 {f2(lin['r_age_Y']['미도입'])}; r(경력,Y) 도입 {f2(lin['r_exp_Y']['도입'])}·미도입 {f2(lin['r_exp_Y']['미도입'])}. "
             f"제곱항 2개 추가 F({lin['quad_df'][0]},{lin['quad_df'][1]})={f2(lin['quad_F'])}, p={fp(lin['quad_p'])}. 도표 figs/resid_fitted_Y_total.png·scatter_*_Y_by_group.png",
             "제곱항 F는 진단용(모형에 넣지 않음)."],
            ["회귀기울기 동질성", "D×나이·D×경력 추가 진단 모형(Type III)",
             f"D×나이 F={f2(slope['D_x_age_F'])}, p={fp(slope['D_x_age_p'])}; D×경력 F={f2(slope['D_x_exp_F'])}, p={fp(slope['D_x_exp_p'])}; "
             f"결합 F({slope['joint_df'][0]},{slope['joint_df'][1]})={f2(slope['joint_F'])}, p={fp(slope['joint_p'])}",
             "비유의를 동질성 확보의 증거로 쓰지 않음. H2의 D×W_c와 구분."],
            ["잔차 정규성", "Shapiro–Wilk, 잔차 왜도·첨도, Q-Q 도표",
             f"H1: W={dg['shapiro_W']:.3f}, p={fp(dg['shapiro_p'])}, 왜도 {f2(dg['resid_skew'])}, 첨도 {f2(dg['resid_kurt'])} (figs/qq_Y_total.png). 하위요인: {sub_norm}. "
             f"H2: W={h2['diag']['shapiro_W']:.3f}, p={fp(h2['diag']['shapiro_p'])}",
             "하위요인 단일 문항(Y1·Y2)은 5점 이산값이라 정규성 위반이 구조적."],
            ["등분산성", "집단별 잔차 Levene(평균·중앙값 중심), Breusch–Pagan",
             f"H1: Levene(평균) F={f2(dg['levene_mean_F'])}, p={fp(dg['levene_mean_p'])}; Levene(중앙값) F={f2(dg['levene_median_F'])}, p={fp(dg['levene_median_p'])}; "
             f"잔차 SD 도입 {f2(dg['resid_sd_by_group']['도입'])}·미도입 {f2(dg['resid_sd_by_group']['미도입'])}; BP LM={f2(dg['bp_LM'])}, p={fp(dg['bp_p'])}. 하위요인 BP: {sub_bp}. H2 BP p={fp(h2['diag']['bp_p'])}",
             "위반 시 대안(강건 SE 등)은 04장에서 [확정 필요] — 본 분석은 사전 지정 OLS만 보고."],
            ["영향점", "Cook's D>4/n, 레버리지>2p/n, |외적 스튜던트화 잔차|>3",
             f"H1(n={dg['n']}, p={dg['n_params']}): Cook's D>{dg['cooks_cut']:.3f} {len(dg['cooks_flag'])}건({', '.join(i for i, _ in dg['cooks_flag'])}); 최대 {dg['cooks_max']:.3f}({dg['cooks_max_id']}); "
             f"레버리지>{dg['lev_cut']:.3f} {len(dg['lev_flag'])}건({', '.join(i for i, _ in dg['lev_flag'])}); |r|>3 {len(dg['rstud_flag'])}건{('(' + ', '.join(f'{i} {v:.2f}' for i, v in dg['rstud_flag']) + ')') if dg['rstud_flag'] else ''}; 레버리지≈1 {len(dg['lev_one'])}건{('(' + ', '.join(dg['lev_one']) + ')') if dg['lev_one'] else ''}",
             "사례 삭제 없음(사전 지정 기준 없음, 04장 [확정 필요])."],
            ["공선성·식별성", "나이–경력 r, 소속×역할 교차, 희소범주, rank(X)·계수 수, VIF",
             f"r(나이,경력)={f2(r_ae)}; rank(X)={h1['rank']} / 계수 {h1['n_coef']}(완전 순위={h1['full_rank']}), 조건수 {h1['cond_number']:.0f}; "
             f"VIF 최대 {vif[vif_max]:.2f}({term_label(vif_max)}); 성별 여성은 미도입에만 6명(도입 0) — 성별 더미는 미도입 내부에서만 식별; 소속×역할 표는 <표 5-4> 아래",
             "순위 결함 없음 → 추정 가능. 단, 성별–집단 부분 교락·희소범주로 해당 더미 계수는 불안정."],
            ["공통 관측범위", "연속값 겹침, 범주별 집단 분포",
             f"나이 공통 {fnum(t54['panel_A']['A3_age']['common_lo'])}∼{fnum(t54['panel_A']['A3_age']['common_hi'])}(범위 밖 도입 {t54['panel_A']['A3_age']['outside_common_adopt']}·미도입 {t54['panel_A']['A3_age']['outside_common_non']}); "
             f"경력 공통 {fnum(t54['panel_A']['A4_exp']['common_lo'])}∼{fnum(t54['panel_A']['A4_exp']['common_hi'])}(범위 밖 도입 {t54['panel_A']['A4_exp']['outside_common_adopt']}·미도입 {t54['panel_A']['A4_exp']['outside_common_non']}); "
             f"빈 범주: {', '.join(sum((t54['panel_B'][v]['empty_cells'] for v in FACTORS), [])) or '없음'}",
             "여성 범주는 도입 집단에 없어 그 범주에서의 집단 비교는 외삽."],
            ["영역별 분포", "하위요인 천장(5점)/바닥(1점) 비율(도입/미도입)", cf,
             "천장 비율이 높은 영역은 집단 차이 추정이 압축될 수 있음."],
            ["H2 선형성·상호작용", "W–Y 상관, 집단별 W 범위·공통범위, figs/scatter_W_Y_by_group.png",
             f"r(W,Y) 도입 {f2(h2['r_W_Y']['도입'])}·미도입 {f2(h2['r_W_Y']['미도입'])}·전체 {f2(h2['r_W_Y']['전체'])}; "
             f"W 범위 도입 {fnum(w_rng['도입'][0])}∼{fnum(w_rng['도입'][1])}, 미도입 {fnum(w_rng['미도입'][0])}∼{fnum(w_rng['미도입'][1])}, 공통 {fnum(c_lo)}∼{fnum(c_hi)}",
             "±1SD 수준이 공통범위 안인지 <표 5-8> 패널 B에 표기."],
        ]
        md.append(md_table(["점검 항목", "진단 자료·방법", "결과", "추론에 대한 판단·후속 처리"], rows5, [":---", ":---", ":---", ":---"]))
        md.append("\n주: 하위요인 모형의 영향점·rank는 <표 5-7> 패널 B와 results.json `table_5_7.*.diag`. 진단 결과를 근거로 변수·사례를 넣거나 빼지 않았다.\n")
        R["table_5_5"] = dict(H1=diag_h1, slope_homogeneity=slope, linearity=lin, vif=vif, ceiling_floor=ceil_floor,
                              H2=h2["diag"])
        gaps["5-5"] = ["관측의 독립성: 현장 ID·담당자 명부가 없어 확인 불가.",
                       "가정 위반·영향점·식별 실패 시 대안 분석 기준이 04장에서 [확정 필요]라 대안 분석은 수행하지 않음."]

        # ================================================= 표 5-6 작성
        e = h1["emm"]["equal"]
        ep = h1["emm"]["proportional"]
        md.append("---\n## <표 5-6> 안전관리 실효성의 공분산분석 결과(H1)\n")
        md.append(f"모형: `{h1['formula']}` (OLS, Type III 제곱합; 기준범주 소속=시공사·성별=남성·역할=공사감독·총공사비=20억 원 미만; D: 도입=1·미도입=0)\n")
        md.append("### 패널 A — H1 완전사례의 원평균과 조정 평균\n")
        rowsA6 = [
            ["도입", h1["n_adopt"], f2(h1["raw"]["도입"]["mean"]), f2(h1["raw"]["도입"]["sd"]), f2(e["도입"]["est"]), f2(e["도입"]["se"]), fci(e["도입"]["ci_lo"], e["도입"]["ci_hi"])],
            ["미도입", h1["n_non"], f2(h1["raw"]["미도입"]["mean"]), f2(h1["raw"]["미도입"]["sd"]), f2(e["미도입"]["est"]), f2(e["미도입"]["se"]), fci(e["미도입"]["ci_lo"], e["미도입"]["ci_hi"])],
            ["도입−미도입", "해당 없음", f2(h1["raw"]["도입"]["mean"] - h1["raw"]["미도입"]["mean"]), "해당 없음", f2(e["차이"]["est"]), f2(e["차이"]["se"]), fci(e["차이"]["ci_lo"], e["차이"]["ci_hi"])],
        ]
        md.append(md_table(["집단", "n(명·개소)", "원평균(M)", "원표준편차(SD)", "조정 평균·차이", "SE", "95% CI"], rowsA6,
                           [":---", "---:", "---:", "---:", "---:", "---:", ":---"]))
        md.append(f"\n주: 조정 평균 = 만 나이 {h1['covariate_ref']['A3_age']:.2f}세·경력 {h1['covariate_ref']['A4_exp']:.2f}년(H1 완전사례 평균), 범주형 통제요인 **동일 가중**(jamovi EMM 기본, 각 요인 수준을 같은 비중으로 평균). "
                  f"참고 — 비례 가중(관측 비율): 도입 {f2(ep['도입']['est'])}(SE {f2(ep['도입']['se'])}), 미도입 {f2(ep['미도입']['est'])}(SE {f2(ep['미도입']['se'])}), 차이 {f2(ep['차이']['est'])}. "
                  f"가산 모형이므로 가중 방식과 무관하게 조정 평균 차이 = 집단 계수 B(검산 {h1['check_emm_diff_equals_B']}). 원평균은 H1 완전사례(n={h1['n']}) 기준이라 <표 5-3>(n=71)과 미도입이 1명 다르다. "
                  "동일 가중은 표본에 없는 구성(성별 남·여 각 50%, 총공사비 4구간 각 25% 등)을 기준으로 삼기 때문에 두 집단의 조정 평균 수준이 "
                  "원평균보다 크게 낮아진다(여성 더미 계수가 음수). 따라서 조정 평균의 '수준'은 기준분포에 따라 달라지며, 해석 대상은 가중 방식과 무관한 차이(B)이다. "
                  "`[확정 필요: 조정 평균 가중 방식 — 04장]`\n")
        md.append("### 패널 B — 집단 차이와 조정변수의 비표준화 계수\n")
        order1 = ["Intercept", "D", "A3_age", "A4_exp"]
        coefs1 = sorted(h1["coefs"], key=lambda c: order1.index(c["term"]) if c["term"] in order1 else 9)
        md.append(md_table(["모형 항", "B", "SE", "95% CI", "t", "df", "p"],
                           [[term_label(c["term"]), f2(c["B"]), f2(c["SE"]), fci(c["ci_lo"], c["ci_hi"]), f2(c["t"]), c["df"], fp(c["p"])] for c in coefs1],
                           [":---", "---:", "---:", ":---", "---:", "---:", "---:"]))
        md.append("\n주: 조정변수 계수는 모형 명세 확인용 보고값이며 해석하지 않는다. 성별 더미(여성)는 미도입 집단 안에서만 식별된다.\n")
        md.append("### 패널 C — H1 집단 차이의 검정과 효과크기\n")
        md.append(md_table(["검정 항", "F", "분자 df", "분모 df", "p", "부분 η²"],
                           [["집단(도입=1·미도입=0)", f2(h1["F"]), h1["df1"], h1["df2"], fp(h1["p"]), f3(h1["partial_eta2"])]],
                           [":---", "---:", "---:", "---:", "---:", "---:"]))
        md.append(f"\n주: 부분 η² = SS집단/(SS집단+SS잔차) = {h1['ss_group']:.3f}/({h1['ss_group']:.3f}+{h1['ss_resid']:.3f}). F = t²(검산 {h1['check_F_equals_t2']}). R² = {h1['r2']:.3f}, 수정 R² = {h1['adj_r2']:.3f}.\n")
        md.append("### 패널 D — 분석 규모와 모형 식별 상태\n")
        md.append(md_table(["n(도입/미도입/전체)", "실제 계수 수", "설계행렬 순위 rank(X)", "잔차 df", "가정 위반·추정 가능 여부"],
                           [[f"{h1['n_adopt']}/{h1['n_non']}/{h1['n']}", h1["n_coef"], h1["rank"], h1["df_resid"],
                             f"완전 순위·추정 가능. 잔차 정규성 SW p={fp(dg['shapiro_p'])}, BP p={fp(dg['bp_p'])}, Levene p={fp(dg['levene_mean_p'])}; Cook's D 기준 초과 {len(dg['cooks_flag'])}건(<표 5-5>)"]],
                           [":---", "---:", "---:", "---:", ":---"]))
        md.append(f"\n주: 유효 표본 71명 중 역할 결측 1명(R030)이 빠져 H1 완전사례는 {h1['n']}명. 잔차 df = n − rank(X) = {h1['n']} − {h1['rank']} = {h1['df_resid']}.\n")
        d_ = h1["diff"]
        md.append(f"**H1 요약:** 도입−미도입 조정 차이 B={f2(d_['B'])}점, SE={f2(d_['SE'])}, 95% CI {fci(d_['ci_lo'], d_['ci_hi'])}, "
                  f"F({h1['df1']},{h1['df2']})={f2(h1['F'])}, p={fp(h1['p'])}, 부분 η²={f3(h1['partial_eta2'])} → 판정: **{h1_judge}** "
                  "(사전 기준: 차이 양수이고 p<.05). 자기보고 집단·Python 산출 기준이며 jamovi 대조 전.\n")
        s_ = sens
        sd_ = s_["diff"]
        md.append("### 사전 지정 민감도 분석 — 시스템 경로 미작동 도입 현장 제외(참고)\n")
        md.append(md_table(["구분", "제외 도입 개소", "n(도입/미도입/전체)", "조정 차이 B", "SE", "95% CI", "F(df)", "p", "부분 η²", "rank(X)/계수 수", "잔차 df"],
                           [["sys_path=0 도입 제외", s_["excluded_adopt_n"], f"{s_['n_adopt']}/{s_['n_non']}/{s_['n']}", f2(sd_["B"]), f2(sd_["SE"]),
                             fci(sd_["ci_lo"], sd_["ci_hi"]), f"{f2(s_['F'])}({s_['df1']},{s_['df2']})", fp(s_["p"]), f3(s_["partial_eta2"]),
                             f"{s_['rank']}/{s_['n_coef']}", s_["df_resid"]]],
                           [":---", "---:", ":---", "---:", "---:", ":---", ":---", "---:", "---:", ":---", "---:"]))
        empties = {FACTORS[k][0]: [cd.CAT_LABELS[FACTORS[k][1]][int(v)] for v in vs] for k, vs in s_["empty_levels_in_adopt"].items() if vs}
        md.append(f"\n주: 주 결과는 전체 표본이며 이 분석은 참고. 남은 도입 사례가 {s_['n_adopt']}명뿐이고, 도입 부분집단에 관측되지 않는 범주: "
                  f"{'; '.join(f'{k} {v}' for k, v in empties.items()) or '없음'}. 따라서 집단 차이는 이 범주들에서 미도입 자료에 의존한 외삽이며 추정이 매우 불안정하다. "
                  "A-7 ⑤(기타) 2건은 판정 기준 `[확정 필요]` 상태에서 미작동으로 제외됐다.\n")
        ps = s_["path_split"]
        md.append("경로 작동/미작동 대비(H1 완전사례, Y_total 원평균): "
                  + "; ".join(f"{lab} n={ps[k]['n']}, M={f2(ps[k]['mean'])}(SD {f2(ps[k]['sd'])})"
                              for k, lab in (("adopt_sys_path_1", "도입 sys_path=1"), ("adopt_sys_path_0", "도입 sys_path=0"),
                                             ("non_adopt", "미도입"))) + ".\n")
        rob = R["table_5_6"]["robustness"]
        md.append("### 패널 E — H1 집단 차이의 사후 강건성 점검\n")
        rowsE6 = []
        for k, lab in (("base", "기본 모형(원고 보고값)"), ("hc3", "HC3 강건 표준오차"), ("cooks_excluded", "Cook's D 초과 4건 제외"),
                       ("straightlining_excluded", "일렬응답 1건 제외"), ("male_only", "여성 제외(남성만)"),
                       ("exclude_cost1", "총공사비 20억 원 미만 제외"), ("exclude_role_other", "역할 기타 제외"),
                       ("unadjusted", "공변량 없는 단순 비교")):
            r_ = rob[k]
            rowsE6.append([lab, r_["n"], f2(r_["B"]), f2(r_["SE"]), fci(r_["ci_lo"], r_["ci_hi"]), fp(r_["p"]),
                           ", ".join(r_["excluded_ids"]) or "없음"])
        md.append(md_table(["점검 조건", "n", "조정 차이 B", "SE", "95% CI", "p", "제외 resp_id"], rowsE6,
                           [":---", "---:", "---:", "---:", ":---", "---:", ":---"]))
        md.append(f"\n주: 사후 점검이며 사례·변수 선택에 사용하지 않는다. 공변량 없는 단순 비교의 Mann–Whitney U={rob['unadjusted']['mann_whitney_U']:.1f}, "
                  f"p={fp(rob['unadjusted']['mann_whitney_p'])}. 일렬응답 = B1∼B8·C1∼C4 12문항이 모두 같은 값"
                  f"({len(rob['straightlining_excluded']['excluded_ids'])}건). Cook's D 제외 사례는 <표 5-5>의 {len(rob['cooks_excluded']['excluded_ids'])}건"
                  f"(일치={rob['cooks_excluded']['matches_table_5_5_flags']}). 부분 표본에서 기준범주가 빠진 요인은 남은 최소 수준을 기준으로 재적합했다(B 불변). "
                  f"HC3는 잔차 df의 t분포로 검정. 검출 가능 최소 효과(80% 검정력) = (1.96+0.84)×SE = {f2(rob['mde80']['value'])}점.\n")
        gaps["5-6"] = ["jamovi 버전·옵션 [확정 필요] — Python 산출값을 jamovi로 재확인해야 함.",
                       "조정 평균 가중 방식 [확정 필요] — 동일 가중을 주값, 비례 가중을 참고로 병기.",
                       "민감도 분석은 도입 n=5라 해석 제한이 크다(결과 숨기지 않고 보고)."]

        # ================================================= 표 5-7 작성
        md.append("---\n## <표 5-7> 하위요인별 공분산분석 결과(Holm 보정 포함)\n")
        md.append("### 패널 A — 영역별 조정 차이와 추론값\n")
        rows7 = []
        for sc in SUBSCALES:
            s = subs[sc]
            eq = s["emm"]["equal"]
            rows7.append([SUBSCALES[sc], f"{s['n_adopt']}/{s['n_non']}", f"{f2(eq['도입']['est'])}({f2(eq['도입']['se'])})",
                          f"{f2(eq['미도입']['est'])}({f2(eq['미도입']['se'])})", f"{f2(s['diff']['B'])}({f2(s['diff']['SE'])})",
                          fci(s["diff"]["ci_lo"], s["diff"]["ci_hi"]), f2(s["F"]), f"{s['df1']}, {s['df2']}", fp(s["p"]), fp(s["p_holm"]),
                          f3(s["partial_eta2"])])
        md.append(md_table(["하위요인", "n(도입/미도입)", "도입 조정 평균(SE)", "미도입 조정 평균(SE)", "차이 B(SE)", "차이 95% CI", "F",
                            "분자·분모 df", "비보정 p", "Holm 보정 p", "부분 η²"], rows7,
                           [":---", ":---", ":---", ":---", ":---", ":---", "---:", ":---", "---:", "---:", "---:"]))
        md.append("\n주: 다섯 영역의 집단 항 p 5개를 하나의 검정군으로 Holm 보정(statsmodels multipletests, method='holm'). CI는 비보정. 조정 평균은 H1과 같은 기준(동일 가중).\n")
        md.append("### 패널 B — 영역별 모형 식별 상태와 가정 점검 요약\n")
        rows7b = []
        for sc in SUBSCALES:
            s = subs[sc]
            dg_ = s["diag"]
            rows7b.append([SUBSCALES[sc], s["n_coef"], s["rank"], s["df_resid"],
                           f"추정 가능(완전 순위={s['full_rank']}). SW p={fp(dg_['shapiro_p'])}, BP p={fp(dg_['bp_p'])}, Levene p={fp(dg_['levene_mean_p'])}, "
                           f"Cook's D 초과 {len(dg_['cooks_flag'])}건, |r|>3 {len(dg_['rstud_flag'])}건"])
        md.append(md_table(["하위요인", "실제 계수 수", "설계행렬 순위 rank(X)", "잔차 df", "가정 위반·추정 가능 여부"], rows7b,
                           [":---", "---:", "---:", "---:", ":---"]))
        md.append("")
        gaps["5-7"] = ["없음. 단일 문항 영역(조기인지·대응 신속성)은 5점 이산 종속변수라 OLS 잔차 정규성 가정이 구조적으로 맞지 않음을 해석에 반영해야 함."]

        # ================================================= 표 5-8 작성
        md.append("---\n## <표 5-8> 집단과 안전관리 예산 제약성의 상호작용 분석 결과(H2)\n")
        md.append(f"모형: H1 모형 + W_c + D×W_c. W_c = W − {w_mean:.4f}(H2 완전사례 n={h2['n']}의 W 평균), W의 SD = {w_sd:.4f}.\n")
        md.append("### 패널 A — 모형 항의 계수와 신뢰구간\n")
        order = ["D", "W_c", "D:W_c", "A3_age", "A4_exp", "Intercept"]
        coefs2 = sorted(h2["coefs"], key=lambda c: order.index(c["term"]) if c["term"] in order else 99)
        md.append(md_table(["모형 항", "B", "SE", "95% CI", "t", "df", "p"],
                           [[term_label(c["term"]), f2(c["B"]), f2(c["SE"]), fci(c["ci_lo"], c["ci_hi"]), f2(c["t"]), c["df"], fp(c["p"])] for c in coefs2],
                           [":---", "---:", "---:", ":---", "---:", "---:", "---:"]))
        sa = h2["slope_adopt"]
        md.append(f"\n주: 집단 계수 = W가 H2 표본 평균일 때의 차이. W_c 계수 = 미도입 집단의 W 기울기(처리 부호화); 도입 집단의 W 기울기 = {f2(sa['est'])}(SE {f2(sa['se'])}, 95% CI {fci(sa['ci_lo'], sa['ci_hi'])}).\n")
        md.append("### 패널 B — 공통 관측범위 안의 조건부 집단 차이\n")
        rows8 = []
        for lab, c in cond.items():
            if c["inside_common"]:
                rows8.append([lab, f"W={c['W']:.2f}(W_c={c['W_c']:+.2f}), 공통범위 안", f2(c["adopt"]["est"]), f2(c["non"]["est"]),
                              f2(c["diff"]["est"]), f2(c["diff"]["se"]), fci(c["diff"]["ci_lo"], c["diff"]["ci_hi"]), f2(c["diff"]["t"]),
                              c["diff"]["df"], fp(c["diff"]["p"])])
            else:
                rows8.append([lab, f"W={c['W']:.2f} — {c['excluded_reason']}"] + ["제외"] * 8)
        md.append(md_table(["예산 제약성 수준", "실제 W값·공통범위 포함 여부", "도입 조정 평균", "미도입 조정 평균", "조건부 차이 B", "SE", "95% CI", "t", "df", "p"],
                           rows8, [":---", ":---", "---:", "---:", "---:", "---:", ":---", "---:", ":---", "---:"]))
        md.append(f"\n주: W 범위 도입 {fnum(w_rng['도입'][0])}∼{fnum(w_rng['도입'][1])}, 미도입 {fnum(w_rng['미도입'][0])}∼{fnum(w_rng['미도입'][1])}, "
                  f"공통 관측범위 {fnum(c_lo)}∼{fnum(c_hi)}, H2 완전사례 SD {w_sd:.2f}. 조건부 차이 = β_D + β_DW·w(검산 일치), SE는 계수 공분산행렬로 산출. "
                  "조정 평균은 H1과 같은 공변량 기준값·범주형 동일 가중. 비례 가중(관측 분포) 조정 평균 도입/미도입: "
                  + "; ".join(f"{lab} {f2(c['adopt_proportional']['est'])}/{f2(c['non_proportional']['est'])}"
                              for lab, c in cond.items() if c["inside_common"]) + "(차이는 가중 방식과 무관).\n")
        md.append("### 패널 C — 상호작용 검정과 H2의 분석 규모·식별 상태\n")
        md.append(md_table(["D×W_c의 F", "분자·분모 df", "p", "n(도입/미도입/전체)", "실제 계수 수", "rank(X)", "잔차 df", "가정 위반·추정 가능 여부"],
                           [[f2(h2["interaction_F"]), f"{h2['interaction_df'][0]}, {h2['interaction_df'][1]}", fp(h2["interaction_p"]),
                             f"{h2['n_adopt']}/{h2['n_non']}/{h2['n']}", h2["n_coef"], h2["rank"], h2["df_resid"],
                             f"추정 가능(완전 순위={h2['full_rank']}). SW p={fp(h2['diag']['shapiro_p'])}, BP p={fp(h2['diag']['bp_p'])}, Cook's D 초과 {len(h2['diag']['cooks_flag'])}건; 상호작용 부분 η²={f3(h2['partial_eta2_interaction'])}"]],
                           ["---:", ":---", "---:", ":---", "---:", "---:", "---:", ":---"]))
        md.append(f"\n주: 검출 가능 최소 상호작용(80% 검정력, 양측 α=.05) = (1.96+0.84)×SE(D×W_c) = 2.80×{f2(h2['mde80']['se_base'])} = {f2(h2['mde80']['value'])}.\n")
        hc = h2["hc3"]
        md.append("\n사후 점검 — HC3 강건 표준오차(잔차 df " + str(hc["df_resid"]) + "의 t분포, 정규 근사 p 병기): "
                  + "; ".join(f"{term_label(t)} B={f2(v['B'])}, SE={f2(v['SE'])}, 95% CI {fci(v['ci_lo'], v['ci_hi'])}, "
                              f"p={fp(v['p'])}(정규 근사 {fp(v['p_normal'])})" for t, v in hc["terms"].items()) + ". "
                  + "조건부 집단 차이(HC3): " + "; ".join(f"W {lab} {f2(v['est'])}점, SE={f2(v['se'])}, 95% CI {fci(v['ci_lo'], v['ci_hi'])}, p={fp(v['p'])}"
                                                      for lab, v in hc["conditional_diff"].items()) + ".\n")
        ir = h2["interaction"]
        h2_dir = "양(+)" if ir["B"] > 0 else "음(−)"
        h2_judge = ("지지(상호작용 양수, p<.05)" if ir["p"] < ALPHA and ir["B"] > 0 else
                    "예상과 반대 방향의 유의한 상호작용" if ir["p"] < ALPHA else "기각(상호작용 미확인)")
        md.append(f"\n**H2 요약:** D×W_c B={f2(ir['B'])}(SE {f2(ir['SE'])}, 95% CI {fci(ir['ci_lo'], ir['ci_hi'])}), "
                  f"F({h2['interaction_df'][0]},{h2['interaction_df'][1]})={f2(h2['interaction_F'])}, p={fp(h2['interaction_p'])} → 추정 방향 {h2_dir}, 판정: **{h2_judge}**. "
                  "H2 예상(W가 클수록 차이 증가)은 양(+)의 상호작용이다. 비유의 결과를 상호작용 부재의 증거로 단정하지 않는다.\n")
        gaps["5-8"] = ["상호작용 검정력 미확인(설계상 [확정 필요] 아님, 원고에 한계로 서술)."]

        # ================================================= 표 5-9
        md.append("---\n## <표 5-9> 연구 가설 검증 결과 종합\n")
        md.append(md_table(["가설", "분석방법", "핵심 추정·불확실성", "검정통계량·자유도·p", "검정결과"], [
            ["H1", "공분산분석(ANCOVA, Type III)", f"조정 차이 B={f2(d_['B'])}, 95% CI {fci(d_['ci_lo'], d_['ci_hi'])}, 부분 η²={f3(h1['partial_eta2'])} (n={h1['n']})",
             f"F({h1['df1']},{h1['df2']})={f2(h1['F'])}, p={fp(h1['p'])}", h1_judge],
            ["H2", "상호작용 일반선형모형", f"D×W_c B={f2(ir['B'])}, 95% CI {fci(ir['ci_lo'], ir['ci_hi'])} (n={h2['n']})",
             f"F({h2['interaction_df'][0]},{h2['interaction_df'][1]})={f2(h2['interaction_F'])}, p={fp(h2['interaction_p'])}", h2_judge],
        ], [":---", ":---", ":---", ":---", ":---"]))
        md.append("\n주: α=.05 양측. 집단은 자기보고 기준(명부 대조 전), Python 산출(jamovi 대조 전). 검정력 충분성 미검증.\n")
        R["table_5_9"] = dict(H1=dict(judgement=h1_judge, B=d_["B"], ci=[d_["ci_lo"], d_["ci_hi"]], F=h1["F"], df=[h1["df1"], h1["df2"]],
                                      p=h1["p"], partial_eta2=h1["partial_eta2"], n=h1["n"]),
                              H2=dict(judgement=h2_judge, B=ir["B"], ci=[ir["ci_lo"], ir["ci_hi"]], F=h2["interaction_F"],
                                      df=list(h2["interaction_df"]), p=h2["interaction_p"], n=h2["n"]))
        gaps["5-9"] = ["검정력 충분성은 미검증(원고 주와 같음)."]

        # ================================================= A-9 조합
        md.append("---\n## A-9 지시 전달 소요시간의 집단 차이와 <표 4-6> 조합(잠정)\n")
        md.append(md_table(["보기", "도입 %", "미도입 %", "도입 누적 %", "미도입 누적 %"],
                           [[f"{'①②③④'[i]} {cd.CAT_LABELS['A9'][i + 1]}", fpct(a9["prop_adopt"][i]), fpct(a9["prop_non"][i]),
                             fpct(a9["cum_adopt"][i]), fpct(a9["cum_non"][i])] for i in range(4)],
                           [":---", "---:", "---:", "---:", "---:"]))
        md.append(f"\nFisher 정확검정 p={fp(a9['p_fisher'])}. 방향: '즉시∼5분'은 미도입이 더 많고({fpct(a9['prop_non'][0])}% vs {fpct(a9['prop_adopt'][0])}%), "
                  f"'30분 이내'(①+②) 누적은 도입이 더 많다({fpct(a9['cum_adopt'][1])}% vs {fpct(a9['cum_non'][1])}%) — 한쪽 방향으로 일관되게 몰리지 않는다. "
                  f"탐색적 부가 분석(설계 외): Mann–Whitney U={a9['exploratory_mann_whitney']['U']:.1f}, p={fp(a9['exploratory_mann_whitney']['p'])}.\n")
        md.append(f"**잠정 판정:** A-9 집단 차이 = {a9['provisional_judgement']}, H1 = {h1_judge} → <표 4-6> 조합 **'{a9['provisional_combination']}'**. "
                  "04장은 A-9 '차이 있음' 판정 기준(보기별 비율 차이 최소 크기·Fisher p 사용 방식)을 `[확정 필요]`로 두고 있으므로, "
                  "여기서는 Fisher p<.05를 임시 기준으로 쓴 잠정 조합이다.\n")
        gaps["A-9"] = ["A-9 '차이 있음' 판정 기준이 04장에서 [확정 필요] — 조합은 잠정."]

        # ================================================= 채울 수 없는 칸
        md.append("---\n## 표별 '채울 수 없는 칸과 이유'\n")
        for k, v in gaps.items():
            md.append(f"**<표 {k}>**" if k[0].isdigit() else f"**{k}**")
            md.extend(f"- {x}" for x in v)
            md.append("")
        R["unfillable"] = gaps

        # ================================================= 검증
        xl = None
        try:
            import openpyxl
            wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
            ws = wb["빈도_기술통계"]
            vals = {}
            for r in range(1, ws.max_row + 1):
                if ws.cell(r, 1).value == "안전관리 실효성 전체":
                    vals["도입"] = ws.cell(r, 4).value
                    vals["미도입"] = ws.cell(r + 1, 4).value
                    break
            xl = dict(excel=vals, python={g: t53["panel_A"]["Y_total"][g]["mean"] for g in ("도입", "미도입")})
            xl["match"] = all(abs(xl["excel"][g] - xl["python"][g]) < 1e-9 for g in ("도입", "미도입"))
        except Exception as ex:  # noqa: BLE001
            xl = dict(error=str(ex), match=False)
        checks = dict(
            y_total_raw_mean_matches_excel=xl,
            emm_diff_equals_B=dict(H1=h1["check_emm_diff_equals_B"], **{sc: subs[sc]["check_emm_diff_equals_B"] for sc in SUBSCALES},
                                   sensitivity=sens["check_emm_diff_equals_B"]),
            F_equals_t2=dict(H1=h1["check_F_equals_t2"], **{sc: subs[sc]["check_F_equals_t2"] for sc in SUBSCALES}),
            holm_ge_raw_all=all(subs[sc]["p_holm"] >= subs[sc]["p"] - 1e-15 for sc in SUBSCALES),
            partial_eta2_in_0_1=all(0 <= x <= 1 for x in [h1["partial_eta2"], sens["partial_eta2"], h2["partial_eta2_interaction"]]
                                    + [subs[sc]["partial_eta2"] for sc in SUBSCALES]),
            full_rank=dict(H1=h1["full_rank"], H2=h2["full_rank"], sensitivity=sens["full_rank"],
                           **{sc: subs[sc]["full_rank"] for sc in SUBSCALES}),
            conditional_diff_formula=all(c.get("check_diff_equals_bD_plus_bDW_w", True) for c in cond.values()),
            fisher_2x2_crosscheck=None,
        )
        # Fisher 열거 구현을 2×2에서 scipy와 대조
        t22 = np.array([[int(((G["도입"].A2_sex) == 1).sum()), int(((G["도입"].A2_sex) == 2).sum())],
                        [int(((G["미도입"].A2_sex) == 1).sum()), int(((G["미도입"].A2_sex) == 2).sum())]])
        mine = fisher_exact_2xk(t22)[0]
        ref22 = float(stats.fisher_exact(t22).pvalue)
        checks["fisher_2x2_crosscheck"] = dict(table=t22.tolist(), mine=mine, scipy=ref22, match=abs(mine - ref22) < 1e-9)
        R["checks"] = checks

    R["warnings"] = [f"{w.category.__name__}: {w.message}" for w in wlist]
    md.append("---\n## 검증 기록\n")
    md.append(f"- Y_total 원평균 ↔ Task A 엑셀 `빈도_기술통계`: {checks['y_total_raw_mean_matches_excel']}")
    md.append(f"- 조정 평균 차이 = 집단 계수 B(동일 가중): {checks['emm_diff_equals_B']}")
    md.append(f"- Holm p ≥ 비보정 p 전건: {checks['holm_ge_raw_all']} · 부분 η² ∈ [0,1]: {checks['partial_eta2_in_0_1']} · F = t²: {all(checks['F_equals_t2'].values())}")
    md.append(f"- 완전 순위: {checks['full_rank']} · 조건부 차이 식 검산: {checks['conditional_diff_formula']} · Fisher 열거 2×2 ↔ scipy: {checks['fisher_2x2_crosscheck']['match']}")
    md.append(f"- 실행 경고 {len(R['warnings'])}건: " + ("; ".join(sorted(set(R["warnings"]))) if R["warnings"] else "없음"))

    md_text = "\n".join(md) + "\n"
    # 표 수 = 머리행 다음 구분행 개수
    n_tables = sum(1 for line in md_text.splitlines() if line.startswith("|") and set(line.replace("|", "").strip()) <= set(":- "))
    R["meta"]["md_tables"] = n_tables
    MD_PATH.write_text(md_text, encoding="utf-8")
    JSON_PATH.write_text(json.dumps(jnum(R), ensure_ascii=False, indent=1), encoding="utf-8")

    def leaves(x):
        if isinstance(x, dict):
            return sum(leaves(v) for v in x.values())
        if isinstance(x, list):
            return sum(leaves(v) for v in x) if x else 1
        return 1

    figs = sorted(p.name for p in FIG_DIR.glob("*.png"))
    summary = dict(md_tables=n_tables, json_top_keys=list(R.keys()), json_leaf_values=leaves(jnum(R)),
                   figs=figs, warnings=len(R["warnings"]), checks=jnum(checks),
                   H1=dict(n=h1["n"], B=d_["B"], p=h1["p"], eta2=h1["partial_eta2"], rank=h1["rank"], n_coef=h1["n_coef"]),
                   H2=dict(n=h2["n"], B=ir["B"], p=h2["interaction_p"], rank=h2["rank"], n_coef=h2["n_coef"]))
    log.write(json.dumps(summary, ensure_ascii=False, indent=1) + "\n")
    log.write("\n".join(R["warnings"]) + "\n")
    LOG_PATH.write_text(log.getvalue(), encoding="utf-8")
    print(log.getvalue())
    ok = (checks["y_total_raw_mean_matches_excel"].get("match") and all(checks["emm_diff_equals_B"].values())
          and checks["holm_ge_raw_all"] and checks["partial_eta2_in_0_1"] and checks["fisher_2x2_crosscheck"]["match"]
          and R["meta"]["source_sha_ok"] and R["meta"]["csv_matches_coding_py"])
    print("CHECKS", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
